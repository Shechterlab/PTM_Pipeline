#!/usr/bin/env python3
"""Turn the download cache of scripts/37 into readable tables under inputs/interactions/.

scripts/37_export_interaction_data.py downloads its sources into a cache (PTM_data/interactions_cache): bulk files,
one API answer per protein or structure, and journal spreadsheets. This script reads that cache once and writes one
plain tab-separated table per source, holding what the source says: accessions, positions, partners, structure
identifiers and scores. It filters only where a source covers far more than the pipeline can use (other species,
unreviewed entries, scores below the lowest floor the pipeline applies); the analysis choices (the floors the website
shows, the sequence checks, which structures to examine) stay in scripts/37.

Alongside the tables it writes
  MANIFEST.tsv         one row per table: rows, MD5, what it holds, source, date, licence, how it was made and the
                       cached files it was made from;
  raw_cache_files.tsv  every cached file with its size, MD5 and file date, so a table can be traced to the download
                       it came from even where the download itself is not kept.

Tables from sources whose licence does not yet allow publishing are written too, marked held_pending_licence in the
MANIFEST and listed in .gitignore, so they stay out of Git until that is settled.

Dates are the cache file dates (the day each download was saved). Conventions: UTF-8, tab-separated, one header row,
lists joined with "; ", booleans true/false, empty for missing.

    python scripts/37a_stage_interactions.py --cache PTM_data/interactions_cache --out inputs/interactions
    python scripts/37a_stage_interactions.py --topup PTM_data/interactions_topup --out inputs/interactions
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import gzip
import hashlib
import importlib.util
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

HERE = Path(__file__).resolve().parent
_spec = importlib.util.spec_from_file_location('export37', HERE / '37_export_interaction_data.py')
x37 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(x37)
_spec = importlib.util.spec_from_file_location('interaction_inputs', HERE / 'interaction_inputs.py')
ii = importlib.util.module_from_spec(_spec)   # the readers of the staged tables
_spec.loader.exec_module(ii)

HEADER_CHECK = re.compile(r'[\t\r\n]')
MANIFEST_COLUMNS = ['file', 'bytes', 'md5', 'rows', 'content', 'source', 'version_or_date', 'citation', 'licence',
                    'how_made', 'derived_from', 'redistribution']


# ----------------------------------------------------------------------------- helpers
def md5_file(path: Path) -> str:
    h = hashlib.md5()
    with open(path, 'rb') as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def file_date(path: Path) -> str:
    return dt.date.fromtimestamp(path.stat().st_mtime).isoformat()


def dates_of(paths) -> str:
    """One date, or the first and last date of a set of files."""
    ds = sorted({file_date(p) for p in paths})
    return '' if not ds else ds[0] if len(ds) == 1 else f'{ds[0]} to {ds[-1]}'


def text(value) -> str:
    """Free text from a source, made safe for one TSV cell: no tabs or line breaks; ';' (the list separator) becomes ','."""
    return re.sub(r'\s+', ' ', str(value or '')).replace(';', ',').strip()


def rel(release: str, date: str) -> str:
    return f'UniProt release {release}, downloaded {date}' if release else date


def cell(value) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, (list, tuple, set)):
        return '; '.join(cell(v) for v in value)
    if isinstance(value, float):
        return f'{value:.6g}'
    s = str(value)
    if HEADER_CHECK.search(s):
        raise ValueError(f'tab or line break in a value: {s[:60]!r}')
    return s


def ranges(positions) -> str:
    """Sorted integer positions as comma-separated runs: 3,4,5,9 -> 3-5,9."""
    runs = []
    for v in sorted(set(positions)):
        if runs and v == runs[-1][1] + 1:
            runs[-1][1] = v
        else:
            runs.append([v, v])
    return ','.join(f'{a}-{b}' if b > a else str(a) for a, b in runs)


class Writer:
    def __init__(self, out: Path, cache: Path, keep_existing: bool = False):
        # keep_existing (top-ups): a dated table already staged is never overwritten; a second top-up of the same day
        # writes NAME_DATE_2.tsv, which sorts after it (scripts/37 reads every table of a kind, oldest first).
        self.out, self.cache, self.manifest, self.summary, self.keep_existing = out, cache, [], {}, keep_existing

    def free(self, name: str) -> str:
        if not self.keep_existing:
            return name
        stem, ext, k = name[:name.rindex('.')], name[name.rindex('.'):], 1
        while (self.out / name).exists():
            k += 1
            name = f'{stem}_{k}{ext}'
        return name

    def table(self, name: str, columns: list[str], rows, *, content: str, source: str, version: str, citation: str,
              licence: str, how: str, derived: list[str], redistribution: str = 'git') -> int:
        path = self.out / self.free(name)
        name = path.name
        n = 0
        with open(path, 'w', encoding='utf-8', newline='\n') as fh:
            fh.write('\t'.join(columns) + '\n')
            for row in rows:
                if len(row) != len(columns):
                    raise ValueError(f'{name}: {len(row)} values for {len(columns)} columns')
                fh.write('\t'.join(cell(v) for v in row) + '\n')
                n += 1
        self._add(path, n, content, source, version, citation, licence, how, derived, redistribution)
        return n

    def fasta(self, name: str, records, *, content: str, source: str, version: str, citation: str, licence: str, how: str,
              derived: list[str], redistribution: str = 'git') -> int:
        """records: (accession, header text, sequence); 60 residues per line."""
        path = self.out / self.free(name)
        n = 0
        with open(path, 'w', encoding='utf-8', newline='\n') as fh:
            for acc, head, seq in records:
                fh.write(f'>{acc} {head}'.rstrip() + '\n')
                for k in range(0, len(seq), 60):
                    fh.write(seq[k:k + 60] + '\n')
                n += 1
        self._add(path, n, content, source, version, citation, licence, how, derived, redistribution)
        return n

    def _add(self, path, n, content, source, version, citation, licence, how, derived, redistribution):
        self.manifest.append([f'interactions/{path.name}', path.stat().st_size, md5_file(path), n, content, source, version,
                              citation, licence, how, '; '.join(derived), redistribution])
        print(f'{path.name}: {n:,} rows, {path.stat().st_size / 1e6:.1f} MB')

    def write_manifest(self):
        with open(self.out / 'MANIFEST.tsv', 'w', encoding='utf-8', newline='\n') as fh:
            fh.write('\t'.join(MANIFEST_COLUMNS) + '\n')
            for row in self.manifest:
                fh.write('\t'.join(cell(v) for v in row) + '\n')


# ----------------------------------------------------------------------------- UniProt
def stage_uniprot(w: Writer, c: Path, release: str, reference: set):
    classes = {}
    for key in ('tudor', 'wd40', 'phd'):
        lines = (c / f'uniprot_human_{key}.tsv').read_text(encoding='utf-8').splitlines()[1:]
        classes[key] = {ln.split('\t')[0].strip() for ln in lines if ln.strip()}
    rows, seqs = [], {}
    with x37.open_text(c / 'uniprot_human_reviewed.tsv.gz') as fh:
        for r in csv.DictReader(fh, delimiter='\t'):
            seq = r.get('Sequence') or ''
            seqs[r['Entry']] = seq
            sid = (r.get('STRING') or '').split(';')[0].strip()
            genes = [g.strip() for g in (r.get('Gene Names (primary)') or '').split(';') if g.strip()]
            rows.append([r['Entry'], genes, sid, len(seq), hashlib.md5(seq.encode()).hexdigest().upper(),
                         r['Entry'] in reference if reference else None,
                         r['Entry'] in classes['tudor'], r['Entry'] in classes['wd40'], r['Entry'] in classes['phd']])
    rows.sort(key=lambda x: x[0])
    d = dates_of([c / 'uniprot_human_reviewed.tsv.gz'] + [c / f'uniprot_human_{k}.tsv' for k in classes])
    w.table(f'uniprot_human_reviewed_{d[:10]}.tsv',
            ['uniprot_accession', 'gene_names_primary', 'string_id', 'sequence_length', 'sequence_md5', 'in_reference_proteome',
             'tudor_domain', 'wd40_repeat', 'phd_finger'], rows,
            content='Reviewed human UniProtKB entries: primary gene name(s), STRING identifier, sequence length and MD5, whether the '
                    'entry is in the reference FASTA (inputs/reference, reference proteome UP000005640), and '
                    'membership of the InterPro Tudor (IPR002999), WD40 repeat (IPR001680) and PHD finger (IPR001965, IPR019787) entries',
            source=f'{x37.UNIPROT_STREAM} (query organism_id:9606 AND reviewed:true; xref:interpro-<entry> for the three classes)',
            version=rel(release, d), citation='UniProt Consortium, Nucleic Acids Res 53:D609 (2025); Blum et al., Nucleic Acids Res 53:D444 (2025)',
            licence='CC BY 4.0', how='scripts/37a_stage_interactions.py stage_uniprot; the sequences are in the FASTA file of the same date',
            derived=['uniprot_human_reviewed.tsv.gz', 'uniprot_human_tudor.tsv', 'uniprot_human_wd40.tsv', 'uniprot_human_phd.tsv'])
    w.fasta(f'uniprot_human_reviewed_{d[:10]}.fasta', [(r[0], f'GN={r[1][0]}' if r[1] else '', seqs[r[0]]) for r in rows],
            content='Canonical sequences of the reviewed human UniProtKB entries in the download scripts/37 used (header: accession and '
                    'primary gene name); scripts/37 checks crosslinks, IntAct features and RBDmap fragments against these sequences',
            source=f'{x37.UNIPROT_STREAM} (query organism_id:9606 AND reviewed:true, field sequence)', version=rel(release, d),
            citation='UniProt Consortium, Nucleic Acids Res 53:D609 (2025)', licence='CC BY 4.0',
            how='scripts/37a_stage_interactions.py stage_uniprot', derived=['uniprot_human_reviewed.tsv.gz'])

    feat = re.compile(r'(REGION|DOMAIN|REPEAT|ZN_FING) [<?]?(\d+)\.\.[>?]?(\d+);(?: /note="([^"]*)")?')
    kind = {'REGION': 'disordered_region', 'DOMAIN': 'domain', 'REPEAT': 'repeat', 'ZN_FING': 'zinc_finger'}
    reg = []
    with x37.open_text(c / 'uniprot_human_regions.tsv.gz') as fh:
        for r in csv.reader(fh, delimiter='\t'):
            if not r or r[0] == 'Entry':
                continue
            for m in feat.finditer('\t'.join(r[1:])):
                k, a, b, note = m.group(1), int(m.group(2)), int(m.group(3)), m.group(4) or ''
                if k == 'REGION' and note != 'Disordered':
                    continue
                reg.append([r[0], kind[k], a, b, text(note)])
    d = dates_of([c / 'uniprot_human_regions.tsv.gz'])
    w.table(f'uniprot_human_regions_{d}.tsv', ['uniprot_accession', 'feature', 'start', 'end', 'note'], reg,
            content='Reviewed human UniProtKB features: disordered regions (Region features noted "Disordered") and domains, repeats and zinc fingers',
            source=f'{x37.UNIPROT_STREAM} (fields ft_region, ft_domain, ft_repeat, ft_zn_fing)', version=rel(release, d),
            citation='UniProt Consortium, Nucleic Acids Res 53:D609 (2025)', licence='CC BY 4.0',
            how='scripts/37a_stage_interactions.py stage_uniprot (uncertain ends "<" "?" ">" read as the given position)',
            derived=['uniprot_human_regions.tsv.gz'])

    dna = []
    path = c / 'crosslinks' / 'uniprot_dna_bind.tsv'
    with open(path, encoding='utf-8') as fh:
        next(fh)
        for line in fh:
            acc, _, feats = line.rstrip('\n').partition('\t')
            for m in re.finditer(r'DNA_BIND (\d+)\.\.(\d+)(?:; /note="([^"]*)")?', feats):
                dna.append([acc, int(m.group(1)), int(m.group(2)), text(m.group(3))])
    d = dates_of([path])
    w.table(f'uniprot_human_dna_binding_regions_{d}.tsv', ['uniprot_accession', 'start', 'end', 'note'], dna,
            content='Reviewed human UniProtKB DNA-binding regions (DNA_BIND features)', source=x37.UNIPROT_DNA_BIND, version=rel(release, d),
            citation='UniProt Consortium, Nucleic Acids Res 53:D609 (2025)', licence='CC BY 4.0',
            how='scripts/37a_stage_interactions.py stage_uniprot', derived=['crosslinks/uniprot_dna_bind.tsv'])
    return seqs, rows


# ----------------------------------------------------------------------------- Gene Ontology
def stage_go(w: Writer, c: Path):
    seen, rows = set(), []
    for key, term in x37.GO_TERMS.items():
        with open(c / f'quickgo_{key}.tsv', encoding='utf-8') as fh:
            for r in csv.reader(fh, delimiter='\t'):
                if not r or r[0].startswith('GENE PRODUCT'):
                    continue
                k = (x37.base_acc(r[1]), key, r[1], r[2], r[-1].strip())
                if k not in seen:
                    seen.add(k)
                    rows.append(list(k))
    rows.sort()
    d = dates_of([c / f'quickgo_{k}.tsv' for k in x37.GO_TERMS])
    w.table(f'quickgo_human_rna_dna_binding_{d[:10]}.tsv',
            ['uniprot_accession', 'binding', 'gene_product_id', 'go_term', 'evidence_code'], rows,
            content='GO annotations of human proteins to RNA binding (GO:0003723) or DNA binding (GO:0003677) or their descendants '
                    '(is_a, part_of, occurs_in), one row per protein, term and evidence code',
            source=f'{x37.QUICKGO_DOWNLOAD} (goUsage descendants, taxonId 9606, geneProductType protein)', version=d,
            citation='Gene Ontology Consortium, Genetics 224:iyad031 (2023); Binns et al., Bioinformatics 25:3045 (2009)',
            licence='CC BY 4.0', how='scripts/37a_stage_interactions.py stage_go (duplicate rows removed)',
            derived=[f'quickgo_{k}.tsv' for k in x37.GO_TERMS])


# ----------------------------------------------------------------------------- STRING
def stage_string(w: Writer, c: Path, uniprot_rows):
    sid_to_acc, gene = {}, {}
    for acc, genes, sid, *_ in uniprot_rows:
        gene[acc] = genes[0] if genes else ''
    # The same rule as scripts/37: a STRING identifier names the first reviewed entry (in UniProt's order) that lists it.
    with x37.open_text(c / 'uniprot_human_reviewed.tsv.gz') as fh:
        for r in csv.DictReader(fh, delimiter='\t'):
            sid = (r.get('STRING') or '').split(';')[0].strip()
            if sid and sid not in sid_to_acc:
                sid_to_acc[sid] = r['Entry']
    edges, unmapped = {}, set()
    path = c / Path(x37.STRING_LINKS).name
    with x37.open_text(path) as fh:
        next(fh)
        for line in fh:
            s1, s2, sc = line.split()
            sc = int(sc)
            if sc < 400:
                continue
            a, b = sid_to_acc.get(s1), sid_to_acc.get(s2)
            if not a or not b:
                unmapped.update(s for s, x in ((s1, a), (s2, b)) if not x)
                continue
            if a == b:
                continue
            key = (a, b) if a < b else (b, a)
            if sc > edges.get(key, 0):
                edges[key] = sc
    rows = [[a, gene.get(a, ''), b, gene.get(b, ''), sc] for (a, b), sc in sorted(edges.items())]
    d = dates_of([path])
    w.table(f'string_v{x37.STRING_VERSION}_human_physical_links_score400.tsv',
            ['protein_a', 'gene_a', 'protein_b', 'gene_b', 'combined_score'], rows,
            content=f'STRING v{x37.STRING_VERSION} human physical-subnetwork links with combined score >= 400 (0-1000) between two '
                    'reviewed UniProt entries, each pair once (protein_a < protein_b)',
            source=x37.STRING_LINKS, version=f'v{x37.STRING_VERSION}, downloaded {d}',
            citation='Szklarczyk et al., Nucleic Acids Res 53:D730 (2025)', licence='CC BY 4.0',
            how=f'scripts/37a_stage_interactions.py stage_string: STRING protein ids mapped to reviewed UniProt entries through UniProt '
                f'xref_string (the first entry listing the id); {len(unmapped):,} STRING ids map to no reviewed entry and their links are left out',
            derived=[path.name, 'uniprot_human_reviewed.tsv.gz'])
    return edges, sid_to_acc


# ----------------------------------------------------------------------------- Complex Portal
def stage_complexes(w: Writer, c: Path):
    rows = []
    paths = []
    for kind, url in x37.COMPLEX_URLS.items():
        path = c / Path(url).name
        paths.append(path)
        with open(path, encoding='utf-8') as fh:
            r = csv.reader(fh, delimiter='\t')
            header = next(r)
            col = header.index('Expanded participant list') if 'Expanded participant list' in header else 4
            for row in r:
                if len(row) <= col:
                    continue
                for k, slot in enumerate(x37.parse_members(row[col] or row[4]), 1):
                    for acc in slot:
                        rows.append([row[0], text(row[1]), kind, k, acc])
    d = dates_of(paths)
    w.table(f'complex_portal_human_complexes_{d[:10]}.tsv',
            ['complex_id', 'complex_name', 'curation', 'member_slot', 'uniprot_accession'], rows,
            content='Complex Portal human complexes, curated and predicted (hu.MAP 3.0): one row per member slot and accession; a slot '
                    'that Complex Portal fills with interchangeable proteins has several rows with the same member_slot',
            source='; '.join(x37.COMPLEX_URLS.values()), version=f'current release, downloaded {d}',
            citation='Meldal et al., Nucleic Acids Res 50:D578 (2022); Fischer et al., Mol Syst Biol 21 (2025)', licence='CC0',
            how='scripts/37a_stage_interactions.py stage_complexes, with scripts/37 parse_members (expanded participant list; '
                'stoichiometry dropped; a protein listed twice in a complex is one member)',
            derived=[p.name for p in paths])


# ----------------------------------------------------------------------------- PDBe-KB and SIFTS
def stage_pdbe(w: Writer, c: Path, reviewed: set):
    stage_pdbe_interfaces(w, c)
    stage_sifts(w, c, reviewed)


def stage_pdbe_interfaces(w: Writer, c: Path):
    folder = c / 'pdbe_interfaces'
    files = sorted(folder.glob('*.json'))
    rows, queries = [], []
    for p in files:
        acc = p.stem
        try:
            d = json.loads(p.read_text() or '{}')
        except json.JSONDecodeError:
            queries.append([acc, 'unreadable', file_date(p)])
            continue
        rec = d.get(acc) if isinstance(d, dict) else None
        data = rec.get('data') if rec else None
        queries.append([acc, 'interfaces' if data else 'none', file_date(p)])
        for part in data or []:
            pacc = part.get('accession') or ''
            for r in part.get('residues', []):
                if r.get('indexType', 'UNIPROT') != 'UNIPROT':
                    continue
                pdbs = sorted({e.get('pdbId') for e in r.get('interactingPDBEntries') or [] if e.get('pdbId')})
                rows.append([acc, int(r['startIndex']), int(r['endIndex']), pacc, text(part.get('name')), pdbs])
    # one row per run of consecutive residues sharing the partner and the PDB entries (the API gives one per residue)
    rows.sort(key=lambda x: (x[0], x[3], x[4], x[5], x[1]))
    merged = []
    for r in rows:
        m = merged[-1] if merged else None
        if m and m[0] == r[0] and m[3:] == r[3:] and r[1] <= m[2] + 1:
            m[2] = max(m[2], r[2])
        else:
            merged.append(list(r))
    rows = sorted(merged, key=lambda x: (x[0], x[1], x[3]))
    d = dates_of(files)
    w.table(f'pdbe_kb_interface_residues_{d[:10]}.tsv',
            ['uniprot_accession', 'start', 'end', 'partner', 'partner_name', 'pdb_entries'], rows,
            content='PDBe-KB interface residues: stretches of a protein (UniProt numbering) that contact a partner in at least one PDB '
                    'structure, with the partner (a UniProt accession, RNA or DNA) and the PDB entries showing the contact',
            source=x37.PDBE_IFACE, version=f'PDBe-KB graph API, answers saved {d}',
            citation='PDBe-KB consortium, Nucleic Acids Res 50:D534 (2022)', licence='EMBL-EBI terms of use (free use with attribution)',
            how='scripts/37a_stage_interactions.py stage_pdbe: every residue of every partner, UniProt numbering only; consecutive '
                'residues with the same partner and PDB entries joined into one start-end row',
            derived=[f'pdbe_interfaces/<accession>.json ({len(files):,} files)'])
    w.table(f'pdbe_kb_interface_queries_{d[:10]}.tsv', ['uniprot_accession', 'answer', 'saved'], queries,
            content='Which proteins were looked up in PDBe-KB and whether the answer listed interfaces (interfaces), none (none) or could '
                    'not be read; a protein not listed was not looked up',
            source=x37.PDBE_IFACE, version=d, citation='PDBe-KB consortium, Nucleic Acids Res 50:D534 (2022)',
            licence='EMBL-EBI terms of use', how='scripts/37a_stage_interactions.py stage_pdbe',
            derived=[f'pdbe_interfaces/<accession>.json ({len(files):,} files)'])


def stage_sifts(w: Writer, c: Path, reviewed: set):
    path = c / Path(x37.SIFTS_OBSERVED).name
    obs = defaultdict(list)
    with x37.open_text(path) as fh:
        header = None
        for r in csv.reader(fh, delimiter='\t'):
            if not r or r[0].startswith('#'):
                continue
            if header is None:
                header = {h.strip(): i for i, h in enumerate(r)}
                continue
            acc = r[header['SP_PRIMARY']]
            if acc in reviewed:
                try:
                    obs[acc].append((int(r[header['SP_BEG']]), int(r[header['SP_END']])))
                except ValueError:
                    pass
    merged = []
    for acc in sorted(obs):
        cur = None
        for a, b in sorted(obs[acc]):
            if cur and a <= cur[1] + 1:
                cur[1] = max(cur[1], b)
            else:
                if cur:
                    merged.append([acc, *cur])
                cur = [a, b]
        merged.append([acc, *cur])
    d = dates_of([path])
    w.table(f'sifts_observed_ranges_human_{d}.tsv', ['uniprot_accession', 'start', 'end'], merged,
            content='Residues of reviewed human proteins observed (modelled) in at least one PDB chain, as merged UniProt-numbered ranges',
            source=x37.SIFTS_OBSERVED, version=f'SIFTS, downloaded {d}', citation='Dana et al., Nucleic Acids Res 47:D482 (2019)',
            licence='EMBL-EBI terms of use', how='scripts/37a_stage_interactions.py stage_pdbe: observed segments of all chains merged per protein',
            derived=[path.name])

    path = c / Path(x37.SIFTS_CHAINS).name
    chains = []
    with x37.open_text(path) as fh:
        header = None
        for r in csv.reader(fh, delimiter='\t'):
            if not r or r[0].startswith('#'):
                continue
            if header is None:
                header = {h.strip(): i for i, h in enumerate(r)}
                continue
            acc = r[header['SP_PRIMARY']]
            beg = r[header['RES_BEG']]
            if x37.base_acc(acc) not in reviewed or not re.fullmatch(r'\d+', beg):
                continue
            try:
                chains.append([r[header['PDB']].lower(), r[header['CHAIN']], acc, int(beg), int(r[header['SP_BEG']]), int(r[header['SP_END']])])
            except ValueError:
                pass
    chains.sort(key=lambda x: (x[0], x[1], x[4]))
    d = dates_of([path])
    w.table(f'sifts_pdb_chain_uniprot_human_{d}.tsv',
            ['pdb_id', 'chain', 'uniprot_accession', 'seqres_start', 'uniprot_start', 'uniprot_end'], chains,
            content='SIFTS mapping of PDB chains to reviewed human UniProt entries (accession as SIFTS gives it, isoform suffix included): '
                    'each segment with its first SEQRES index '
                    '(mmCIF label_seq_id) and UniProt range, used to find a UniProt position in a structure',
            source=x37.SIFTS_CHAINS, version=f'SIFTS, downloaded {d}', citation='Dana et al., Nucleic Acids Res 47:D482 (2019)',
            licence='EMBL-EBI terms of use', how='scripts/37a_stage_interactions.py stage_pdbe: segments with a numeric SEQRES start',
            derived=[path.name])


def stage_modelled(w: Writer, c: Path):
    folder = c / 'modelled_methyl_r'
    entries = sorted(x['identifier'].lower() for x in json.loads((folder / 'entries.json').read_text()).get('result_set', []))
    rows = []
    for pdb in entries:
        try:
            recs = json.loads((folder / f'{pdb}.json').read_text()).get(pdb, [])
        except (OSError, json.JSONDecodeError):
            continue
        for r in recs:
            if r.get('chem_comp_id') in x37.METHYL_R and r.get('residue_number') is not None:
                rows.append([pdb, r.get('chain_id'), int(r['residue_number']), r.get('author_residue_number'), r['chem_comp_id']])
    files = sorted(folder.glob('*.json'))
    d = dates_of(files)
    w.table(f'pdb_methylarginine_residues_{d[:10]}.tsv', ['pdb_id', 'chain', 'seqres_index', 'author_residue_number', 'chem_comp'], rows,
            content='Methylated arginine residues modelled in PDB entries: chemical components 2MR (dimethyl, symmetric), DA2 (dimethyl, '
                    'asymmetric) and NMM (monomethyl); AGM (5-methylarginine, C-methylated) is not counted',
            source=f'{x37.RCSB_SEARCH} (entries with these components); {x37.PDBE_MODRES}', version=d,
            citation='Burley et al., Nucleic Acids Res 51:D488 (2023)', licence='CC0 (PDB data)',
            how='scripts/37a_stage_interactions.py stage_modelled', derived=[f'modelled_methyl_r/ ({len(files)} files)'])


def stage_contacts(w: Writer, c: Path):
    folder = c / 'contacts_seqres'
    files = sorted(folder.glob('*.json'))
    rows = []
    for p in files:
        stem, label = p.stem.rsplit('_', 1)
        pdb, chain = stem.split('_', 1)
        try:
            d = json.loads(p.read_text())
        except json.JSONDecodeError:
            continue
        hits = d.get('hits') or []
        if not hits:
            rows.append([pdb, chain, int(label), d.get('resn', ''), '', '', '', '', ''])
        for ch, lab, seq, comp, dist in hits:
            rows.append([pdb, chain, int(label), d.get('resn', ''), ch, lab, seq, comp, dist])
    rows.sort(key=lambda x: (x[0], x[1], x[2], x[8] if x[8] != '' else 0))
    d = dates_of(files)
    w.table(f'pdb_arginine_contacts_{d[:10]}.tsv',
            ['pdb_id', 'chain', 'seqres_index', 'residue', 'contact_chain', 'contact_seqres_index', 'contact_author_number',
             'contact_residue', 'distance_a'], rows,
            content='Residues of other chains within 4 Å of the side chain (CD, NE, CZ, NH1, NH2) of an arginine at a PDB interface, in the '
                    'first model; a row with empty contact columns is an arginine that was examined and touches nothing within 4 Å',
            source=x37.MODELSERVER, version=f'RCSB ModelServer, read {d}', citation='Burley et al., Nucleic Acids Res 51:D488 (2023)',
            licence='CC0 (PDB data)', how='scripts/37 site_contacts (distances from the residue neighbourhoods it read); collected here',
            derived=[f'contacts_seqres/<pdb>_<chain>_<seqres index>.json ({len(files):,} files)'])


def stage_nucleic_chains(w: Writer, c: Path):
    """The RNA and DNA chains, with their SEQRES sequences, of the PDB entries where a methylarginine meets nucleic acid."""
    files = sorted((c / 'pdbe_molecules').glob('*.json'))
    rows = []
    for p in files:
        try:
            chains = x37.nucleic_chain_rows(json.loads(p.read_text() or '{}').get(p.stem, []))
        except json.JSONDecodeError:
            continue
        rows += [[p.stem, ch, entity, kind, name, organism, prep, histones, len(seq), seq] for ch, entity, kind, name, organism, seq, prep, histones in chains]
        if not chains:
            rows.append([p.stem, '', '', '', '', '', '', '', '', ''])
    rows.sort(key=lambda x: (x[0], x[1]))
    d = dates_of(files)
    w.table(f'pdb_nucleic_acid_chains_{d[:10]}.tsv',
            ['pdb_id', 'chain', 'entity_id', 'nucleic_acid', 'molecule_name', 'source_organism', 'sample_preparation', 'entry_core_histones', 'length', 'sequence'], rows,
            content='The RNA and DNA chains of the PDB entries where a methylated protein touches nucleic acid: molecule type (rna, dna or '
                    'hybrid), name, source organism, sample preparation (Synthetically obtained for a made oligonucleotide), whether the '
                    'entry holds a core histone (H2A, H2B, H3, H4 or a variant: a nucleosome\'s DNA) and SEQRES '
                    'sequence in one-letter codes (modified nucleotides as their parent); a row with an empty chain is an entry looked up '
                    'that has none',
            source=x37.PDBE_MOLECULES, version=f'PDBe API, read {d}', citation='Burley et al., Nucleic Acids Res 51:D488 (2023)',
            licence='CC0 (PDB data)', how='scripts/37 load_nucleic_chains; collected here',
            derived=[f'pdbe_molecules/<pdb>.json ({len(files):,} files)'])


def stage_na_arginines(w: Writer, c: Path):
    """Every arginine modelled in the protein chains read against RNA and DNA (every entry with the protein and nucleic
    acid), and the distance of each side chain to the RNA and DNA chains within 8 Å."""
    files = sorted((c / 'na_arginines').glob('*.json'))
    chains, dists = [], []
    for p in files:
        pdb, chain = p.stem.split('_', 1)
        try:
            res = json.loads(p.read_text())
        except json.JSONDecodeError:
            continue
        chains.append([pdb, chain, len(res['args']), ['%d:%s:%s' % (lab, auth, comp) for lab, auth, comp in res['args']]])
        dists += [[pdb, chain, lab, nch, k, comp, d, within] for lab, nch, k, comp, d, within in res['hits']]
    chains.sort(key=lambda x: (x[0], x[1]))
    dists.sort(key=lambda x: (x[0], x[1], x[2], x[3]))
    d = dates_of(files)
    common = dict(source=x37.MODELSERVER_INTERACTION, version=f'RCSB ModelServer, read {d}', citation='Burley et al., Nucleic Acids Res 51:D488 (2023)',
                  licence='CC0 (PDB data)', how='scripts/37 chain_na_proximity; collected here',
                  derived=[f'na_arginines/<pdb>_<chain>.json ({len(files):,} files)'])
    w.table(f'pdb_chain_arginines_{d[:10]}.tsv', ['pdb_id', 'chain', 'n_arginines', 'arginines'], chains,
            content='The arginines modelled (first model) in each protein chain read against RNA and DNA: every chain of a '
                    'methylated protein (up to two copies) in every PDB entry where PDBe-KB lists the protein with RNA or DNA, '
                    'as SEQRES index:author number:residue (methylated arginines as 2MR, DA2, NMM)', **common)
    w.table(f'pdb_arginine_nucleic_acid_distances_{d[:10]}.tsv',
            ['pdb_id', 'chain', 'seqres_index', 'na_chain', 'closest_na_seqres_index', 'closest_nucleotide', 'side_chain_distance_a', 'na_seqres_within_8a'], dists,
            content='For each arginine of those chains with a side-chain atom (CD, NE, CZ, NH1, NH2) within 8 Å of an RNA or DNA '
                    'chain of the entry: the closest nucleotide, its distance and every nucleotide of that chain within 8 Å', **common)


def stage_na_regions(w: Writer, c: Path):
    """Every residue of a methylated protein's chain within 4 Å of RNA or DNA, one entry per distinct complex."""
    files = sorted((c / 'na_regions').glob('*.json'))
    rows = []
    for p in files:
        pdb, chain = p.stem.split('_', 1)
        try:
            hits = json.loads(p.read_text()).get('hits') or []
        except json.JSONDecodeError:
            continue
        rows += [[pdb, chain, *h] for h in hits] or [[pdb, chain, '', '', '', '', '', '', '']]
    rows.sort(key=lambda x: (x[0], x[1], x[2] if x[2] != '' else 0, x[5] if x[5] != '' else '', x[6] if x[6] != '' else 0))
    d = dates_of(files)
    w.table(f'pdb_nucleic_acid_residue_contacts_{d[:10]}.tsv',
            ['pdb_id', 'chain', 'seqres_index', 'author_number', 'residue', 'na_chain', 'na_seqres_index', 'nucleotide', 'distance_a'], rows,
            content='Residues of a methylated protein\'s chain within 4 Å (any heavy atoms) of an RNA or DNA chain of the entry, in the '
                    'asymmetric unit of the first model, read in one entry per distinct complex (the protein with the same bound '
                    'sequences); a row with empty contact columns is a chain read that touches none',
            source=x37.MODELSERVER_INTERACTION, version=f'RCSB ModelServer, read {d}', citation='Burley et al., Nucleic Acids Res 51:D488 (2023)',
            licence='CC0 (PDB data)', how='scripts/37 chain_na_contacts; collected here',
            derived=[f'na_regions/<pdb>_<chain>.json ({len(files):,} files)'])


def _env(p: str):
    path = Path(p)
    stem, label = path.name[:-len('.json.gz')].rsplit('_', 1)
    pdb, chain = stem.split('_', 1)
    try:
        rows = json.loads(gzip.decompress(path.read_bytes()))
    except (OSError, json.JSONDecodeError):
        return None
    res = x37.analyze_environment(rows, chain, int(label)) if rows else None
    if not res:
        return [pdb, chain, int(label), '', '', '', '', '', '', '', '', '']
    return [pdb, chain, int(label), res['resn'], res['levy'], res['rasa_mono'], res['rasa_cplx'], res['d_na'], res['d_prot'],
            res['na_type'], res['base_hb'], res['n_prot_hb']]


def stage_environments(w: Writer, c: Path, workers: int):
    files = sorted((c / 'environments').glob('*.json.gz'))
    with ProcessPoolExecutor(max_workers=workers) as pool:
        rows = [r for r in pool.map(_env, [str(p) for p in files], chunksize=64) if r]
    rows.sort(key=lambda x: (x[0], x[1], x[2]))
    d = dates_of(files)
    w.table(f'pdb_arginine_interface_environments_{d[:10]}.tsv',
            ['pdb_id', 'chain', 'seqres_index', 'residue', 'levy_class', 'relative_asa_alone', 'relative_asa_in_complex',
             'asa_buried_by_nucleic_acid_a2', 'asa_buried_by_protein_a2', 'nucleic_acid_contact', 'base_hbond_partners',
             'protein_hbonds'], rows,
            content='Geometry of arginines at PDB interfaces, from all atoms within 10 Å (first model): relative solvent accessibility of '
                    'the residue alone and in the complex (Shrake-Rupley, Arg maximum 274 Å²), Levy (2010) class (core, rim, support or '
                    'none), surface buried by nucleic acid and by protein, the kind of nucleic-acid contact (base or backbone hydrogen '
                    'bond, contact, none), the bases hydrogen-bonded and the number of protein hydrogen bonds; empty measures: the '
                    'residue at that index is not an arginine',
            source=x37.MODELSERVER + ' (radius 10 Å)', version=f'RCSB ModelServer, read {d}',
            citation='Levy, J Mol Biol 403:660 (2010); Shrake & Rupley, J Mol Biol 79:351 (1973); Tien et al., PLoS ONE 8:e80635 (2013)',
            licence='CC0 (PDB data)', how='scripts/37 analyze_environment on each saved neighbourhood (the raw atoms are not kept in Git)',
            derived=[f'environments/<pdb>_<chain>_<seqres index>.json.gz ({len(files):,} files)'])


# ----------------------------------------------------------------------------- predicted complexes
def stage_protvar(w: Writer, c: Path):
    for url, name, what in ((x37.PROTVAR_IFACE, 'protvar_predicted_interfaces_2024.05.28_pdockq0.23.tsv',
                             'ProtVar interface residues (within 5 Å) of AlphaFold-Multimer models of human protein pairs, pDockQ >= 0.23'),
                            (x37.PROTVAR_SERVED, 'protvar_served_models_af2complexes_v1.tsv',
                             'Human protein pairs whose AlphaFold-Multimer model the ProtVar API serves, with chains, pDockQ and interface residues')):
        path = c / Path(url).name
        rows, header = [], None
        with x37.open_text(path) as fh:
            r = csv.reader(fh, delimiter='\t')
            header = next(r)
            ix = {h: k for k, h in enumerate(header)}
            for row in r:
                if name.startswith('protvar_predicted'):
                    try:
                        if float(row[ix['pdockq']]) < x37.PRED_TIERS[0]:
                            continue
                    except (ValueError, IndexError):
                        continue
                if name.startswith('protvar_served'):
                    # numbers only in this file: written as runs (3-37,39,43) to keep the table small
                    row = [ranges(int(x) for x in v.split(',') if x.isdigit()) if k in (ix['ifresid1'], ix['ifresid2']) else v
                           for k, v in enumerate(row)]
                rows.append([text(v) if k not in (ix['ifresid1'], ix['ifresid2']) else v for k, v in enumerate(row)])
        served = name.startswith('protvar_served')
        cols = [h + '_ranges' if served and h in ('ifresid1', 'ifresid2') else h for h in header]
        d = dates_of([path])
        w.table(name, cols, rows, content=what + ('; interface residue numbers written as comma-separated runs (3-37,39,43)' if served
                                                 else '; residue lists are comma-separated, as ProtVar gives them'),
                source=url, version=f'{"2024.05.28_interface_summary_5A" if not served else "af2complexes_interfaces_v1"}, downloaded {d}',
                citation='Stephenson et al., Nucleic Acids Res 52:W140 (2024); Burke et al., Nat Struct Mol Biol 30:216 (2023)',
                licence='CC BY 4.0', how='scripts/37a_stage_interactions.py stage_protvar (a copy of the file' +
                (', rows with pDockQ < 0.23 left out)' if not served else '; residue numbers joined into runs)'),
                derived=[path.name], redistribution='git')


def stage_model_contacts(w: Writer, c: Path, served: Path | None = None, measured: str = 'the methylarginines of the published union'):
    """served: the ProtVar served-models list (default: the raw download in the cache, else a staged table with the same
    uniprot_id1, uniprot_id2, chain1 and chain2 columns); measured: whose positions were measured, for the MANIFEST."""
    flags = lambda f: [bool(f & x37.G_CLASH), bool(f & x37.G_SB), bool(f & x37.G_HB), bool(f & x37.G_CATPI)]
    rows, files = [], []
    # ProtVar: the chain of each accession comes from the served-models table, as in scripts/37.
    chain_of = {}
    with x37.open_text(served or c / Path(x37.PROTVAR_SERVED).name) as fh:
        r = csv.reader(fh, delimiter='\t')
        ix = {h: k for k, h in enumerate(next(r))}
        for row in r:
            a, b = row[ix['uniprot_id1']], row[ix['uniprot_id2']]
            chain_of[(a, b)] = {row[ix['chain1']] or 'A': (a, b), row[ix['chain2']] or 'B': (b, a)}
    not_served = 0
    for p in sorted((c / 'protvar_measure').glob('*.json')):
        files.append(p)
        a, b = p.stem.split('_', 1)
        rec = json.loads(p.read_text())
        if not rec.get('served'):
            not_served += 1
        for ch, sites in rec.get('chains', {}).items():
            acc, other = chain_of.get((a, b), {}).get(ch, ('', ''))
            for pos, m in sorted(sites.items(), key=lambda kv: int(kv[0])):
                rows.append(['ProtVar', f'{a}_{b}', acc, ch, int(pos), other, m[0], *flags(m[1]), m[2], m[3],
                             [f'{x}:{lab}' for x, lab in m[4]]])
    # AlphaFold DB: the chain of each accession from the model's metadata in the cache (afdb_prediction), else from the
    # heterodimer tables already staged, with the measured chains' own accessions (positions) over both.
    staged = ii.afdb_heterodimers(ii.Tables(w.out)) if any(w.out.glob('alphafold_db_human_heterodimers_*.tsv')) else {}
    for p in sorted((c / 'afdb_measure').glob('*.json')):
        files.append(p)
        rec = json.loads(p.read_text())
        pf = c / 'afdb_prediction' / f'{p.stem}.json'
        meta = [[m.get('chainId'), m.get('uniprotAccession')] for m in (ii.read_json(pf, []) if pf.exists() else []) if isinstance(m, dict)]
        acc_of = {**dict(zip(staged.get(p.stem, {}).get('chains', []), staged.get(p.stem, {}).get('accessions', []))),
                  **{ch: acc for ch, acc in rec.get('meta', []) or meta}, **{ch: v[0] for ch, v in rec.get('positions', {}).items()}}
        for ch, sites in rec.get('chains', {}).items():
            others = [acc for c2, acc in acc_of.items() if c2 != ch]
            for pos, m in sorted(sites.items(), key=lambda kv: int(kv[0])):
                rows.append(['AlphaFold DB', p.stem, acc_of.get(ch, ''), ch, int(pos), others[0] if others else '', m[0], *flags(m[1]),
                             m[2], m[3], [f'{x}:{lab}' for x, lab in m[4]]])
    d = dates_of(files)
    w.table(f'predicted_model_arginine_contacts_{d[:10]}.tsv',
            ['source', 'model', 'uniprot_accession', 'chain', 'position', 'partner_accession', 'min_distance_a', 'clash',
             'salt_bridge', 'hydrogen_bond', 'cation_pi', 'plddt_arginine', 'plddt_partner_min', 'partner_residues'], rows,
            content='Methylarginines measured in predicted complexes (ProtVar served AlphaFold-Multimer models with pDockQ >= 0.5, and '
                    'AlphaFold DB human heterodimers): arginines within 5 Å of the partner, the closest distance, bonds (salt bridge to '
                    'Asp/Glu carboxylate O <= 4 Å, hydrogen bond to an acceptor O or His ND1/NE2 <= 3.5 Å, cation-pi to an aromatic ring '
                    '<= 5 Å), a clash (< 2.2 Å), pLDDT, and the partner residues named as chain:residue:bonds; measured at the positions of '
                    f'{measured}; {not_served:,} ProtVar pairs listed as served returned no model',
            source=f'{x37.PROTVAR_MODEL}; {x37.AFDB_PREDICTION}', version=f'models read {d}',
            citation='Stephenson et al., Nucleic Acids Res 52:W140 (2024); Han et al., bioRxiv (2026), doi:10.64898/2026.03.27.714458',
            licence='CC BY 4.0', how='scripts/37 measure_protvar_models and measure_afdb_complexes (measure_methyl_r); collected here',
            derived=[f'protvar_measure/ and afdb_measure/ ({len(files):,} files)'])

    lists = sorted((c / 'afdb_complex').glob('*.json'))
    if not lists:
        return
    het, queries = {}, []
    for p in lists:
        try:
            recs = json.loads(p.read_text())
        except json.JSONDecodeError:
            continue
        n = 0
        for x in recs if isinstance(recs, list) else []:
            if x.get('assemblyType') == 'Hetero' and set(x.get('taxId') or []) == {9606} and x.get('modelEntityId'):
                het[x['modelEntityId']] = x
                n += 1
        queries.append([p.stem, n, file_date(p)])
    rows = []
    for mid, x in sorted(het.items()):
        meta = []
        pf = c / 'afdb_prediction' / f'{mid}.json'
        if pf.exists():
            try:
                meta = [[m.get('chainId'), m.get('uniprotAccession')] for m in json.loads(pf.read_text())]
            except (json.JSONDecodeError, AttributeError):
                meta = []
        listed = x.get('uniprotAccession') or x.get('uniprotAccessions') or []   # the API's key is uniprotAccession, a list
        accs = [a for _, a in meta] or ([listed] if isinstance(listed, str) else list(listed))
        rows.append([mid, accs, [ch for ch, _ in meta], x.get('complexPredictionAccuracy_ipTM'), x.get('complexPredictionAccuracy_pDockQ')])
    d = dates_of(lists)
    w.table(f'alphafold_db_human_heterodimers_{d[:10]}.tsv', ['model_id', 'uniprot_accessions', 'chains', 'iptm', 'pdockq'], rows,
            content='AlphaFold DB human heterodimer models that include a protein of the published union, with their chains (where the '
                    'model was read) and the interface scores ipTM and pDockQ',
            source=x37.AFDB_COMPLEX, version=f'AlphaFold DB API, answers saved {d}',
            citation='Han et al., bioRxiv (2026), doi:10.64898/2026.03.27.714458', licence='CC BY 4.0',
            how='scripts/37a_stage_interactions.py stage_model_contacts', derived=[f'afdb_complex/ and afdb_prediction/ ({len(lists):,} lists)'])
    w.table(f'alphafold_db_complex_queries_{d[:10]}.tsv', ['uniprot_accession', 'human_heterodimers', 'saved'], queries,
            content='Which proteins were looked up in the AlphaFold DB complex API and how many human heterodimer models each answer listed',
            source=x37.AFDB_COMPLEX, version=d, citation='Han et al., bioRxiv (2026)', licence='CC BY 4.0',
            how='scripts/37a_stage_interactions.py stage_model_contacts', derived=[f'afdb_complex/ ({len(lists):,} files)'])


# ----------------------------------------------------------------------------- IntAct
def intact_check(kind: int, ranges_: str, original: str, seq: str) -> str:
    """The test scripts/37 applies before using a feature: its ranges must be determined and its curated sequence must be
    the canonical sequence at those ranges (and an arginine, for a methylated-arginine feature)."""
    parts = ranges_.split(',')
    spans = [tuple(map(int, r.split('-'))) for r in parts if re.fullmatch(r'\d+-\d+', r)]
    if len(spans) != len(parts):
        return 'undetermined range'
    if any(a < 1 or b > len(seq) for a, b in spans):
        return 'outside the sequence'
    given = original.split(',') if original not in ('', '-') else []
    if any(k < len(given) and seq[a - 1:b] != given[k] for k, (a, b) in enumerate(spans)):
        return 'mismatch'
    if kind == 2 and any(seq[a - 1:b] != 'R' for a, b in spans):
        return 'not arginine'
    return 'matches' if given else 'no curated sequence'


def stage_intact(w: Writer, c: Path, seqs: dict):
    release = ''
    rp = c / 'intact' / 'release_time.txt'
    if rp.exists():
        release = rp.read_text(encoding='utf-8').strip()
    m = re.search(r'finished on (.+)$', release)
    ver = (m.group(1) if m else release) or 'current'
    kinds = ((0, 'mutations', 'mutations', 'mutations with their effect on the interaction'),
             (1, 'bindings_regions', 'binding_regions', 'binding regions (sufficient, necessary, binding-associated, direct)'),
             (2, 'ptms', 'methylarginines', 'methylated arginines'))
    for kind, name, label, what in kinds:
        path = c / 'intact' / f'{name}.tsv'
        rows, dropped = [], Counter()
        with open(path, encoding='utf-8') as fh:
            next(fh)
            for line in fh:
                x = line.rstrip('\n').split('\t')
                if len(x) < 16 or 'taxid:9606' not in x[10] or not x[7].startswith('uniprotkb:'):
                    continue
                if kind == 2:
                    lab = x[5].lower()
                    if 'arginine' not in lab or 'methyl' not in lab:
                        continue
                else:
                    mi = re.search(r'MI:\d+', x[5])
                    code = (x37.INTACT_EFFECT_MI.get(mi.group(0), 5) if kind == 0 else x37.INTACT_REGION_MI.get(mi.group(0))) if mi else None
                    if code is None:
                        continue
                acc = x[7][len('uniprotkb:'):]
                if acc not in seqs:
                    dropped['isoform or chain' if '-' in acc else 'not a reviewed human entry'] += 1
                    continue
                check = intact_check(kind, x[2], x[3], seqs[acc])
                mt = re.match(r'psi-mi:"?(MI:\d+)"?\((.*)\)$', x[5])
                parts = x37.INTACT_PARTNER.findall(x[11])
                partners = [f'{db}:{ident.strip(chr(34))} ({ptype})' for db, ident, _, ptype in parts]
                pmids = list(dict.fromkeys(m2.group(1) for m2 in re.finditer(r'pubmed:(\d+)', x[12])))   # IntAct's order
                rows.append([x[0], mt.group(1) if mt else '', text(mt.group(2) if mt else x[5]), acc, x[2], check,
                             '' if check in ('matches', 'no curated sequence') else text(x[3]), text(x[4]), partners, pmids, x[14]])
        d = dates_of([path])
        checks = Counter(r[5] for r in rows)
        w.table(f'intact_{label}_human_{d}.tsv',
                ['feature_ac', 'feature_type_mi', 'feature_type', 'uniprot_accession', 'ranges', 'sequence_check',
                 'curated_sequence', 'resulting_sequence', 'interaction_participants', 'pubmed', 'interaction_ac'], rows,
                content=f'IntAct curated features of reviewed human interaction participants: {what}; participants are '
                        'database:identifier (type); sequence_check compares the curated sequence with the canonical UniProt sequence ('
                        + ', '.join(f'{k} {v:,}' for k, v in checks.most_common()) + '); curated_sequence is given only where it '
                        'does not match (where it matches it is the canonical sequence at the range)',
                source=x37.INTACT_FEATURES.format(name=name), version=f'IntAct release finished {ver}; downloaded {d}',
                citation='del Toro et al., Nucleic Acids Res 50:D648 (2022)', licence='CC BY 4.0',
                how='scripts/37a_stage_interactions.py stage_intact: human (taxid 9606) UniProt participants, the feature types '
                    'scripts/37 reads, reviewed canonical accessions only (' + (', '.join(f'{k} {v:,}' for k, v in sorted(dropped.items())) or 'none')
                    + ' left out); scripts/37 uses the rows whose check is "matches" or "no curated sequence"',
                derived=[f'intact/{path.name}', 'intact/release_time.txt'])


# ----------------------------------------------------------------------------- crosslinks
def stage_crosslinks(w: Writer, c: Path):
    import openpyxl
    folder = c / 'crosslinks'
    zpath = folder / 'pepr_media-1.zip'
    s2, s3 = [], []
    with zipfile.ZipFile(zpath) as zf:
        names = {Path(n).name: n for n in zf.namelist()}
        wb = openpyxl.load_workbook(io.BytesIO(zf.read(names['Table_S2.xlsx'])), read_only=True)
        for sheet, label in x37.RNA_XL_SETS:
            rows = wb[sheet].iter_rows(values_only=True)
            ix = {h: k for k, h in enumerate(next(rows))}
            for r in rows:
                pid = str(r[ix['Protein ID']] or '')
                mod = str(r[ix['Modified Sequence']] or '')
                m = re.search(r'([A-Z])\[(\d+\.\d+)\]', re.sub(r'C\[57\.02\d*\]|M\[15\.99\d*\]', lambda x: x.group(0)[0], mod))
                try:
                    pos = int(r[ix['Position']])
                except (TypeError, ValueError):
                    pos = ''
                s2.append([label, x37.base_acc(pid.split(';')[0]), [p.strip() for p in pid.split(';') if p.strip()], pos,
                           m.group(1) if m else '', str(r[ix['Peptide Sequence']] or ''), mod, str(r[ix['RNA Mass']] or '')])
        wb = openpyxl.load_workbook(io.BytesIO(zf.read(names['Table_S3.xlsx'])), read_only=True)
        rows = wb['Crosslinking site'].iter_rows(values_only=True)
        ix = {h: k for k, h in enumerate(next(rows))}
        for r in rows:
            aas = json.loads(r[ix['Focused Site AA']] or '[]')
            poss = json.loads(r[ix['Focused Site Position']] or '[]')
            s3.append([x37.base_acc(str(r[ix['Protein ID']])), str(r[ix['Protein ID']]), aas, poss])
    d = dates_of([zpath])
    common = dict(source=x37.PEPR_ZIP, version=f'bioRxiv v2 (2026) supplement, downloaded {d}',
                  citation='Sha, Kuster & Trendel, bioRxiv (2025), doi:10.1101/2025.09.05.674461; re-analysed: Kramer et al., Nat '
                           'Methods 11:1064 (2014); Trendel et al., Cell 176:391 (2019); Bae et al., Nat Struct Mol Biol 27:678 (2020); '
                           'Bae et al., Nat Commun 12:6026 (2021)', licence='CC BY-NC', derived=['crosslinks/pepr_media-1.zip'])
    w.table('pepr_ms_v2_table_s2_rna_crosslinked_peptides.tsv',
            ['dataset', 'uniprot_accession', 'protein_ids', 'position', 'crosslinked_residue', 'peptide', 'modified_sequence', 'rna_mass'], s2,
            content='pepR-MS compendium Table S2: RNA-crosslinked peptide-spectrum matches of five datasets re-analysed with one pipeline, '
                    'with the residue carrying the nucleotide mass (read from the modified sequence) and its position as the table gives it',
            how='scripts/37a_stage_interactions.py stage_crosslinks (one row per table row; the check against the canonical sequence is '
                'left to scripts/37)', **common)
    w.table('pepr_ms_v2_table_s3_focused_sites.tsv', ['uniprot_accession', 'protein_id', 'focused_residues', 'focused_positions'], s3,
            content='pepR-MS compendium Table S3 (sheet Crosslinking site): the focused crosslinking site(s) of each crosslink',
            how='scripts/37a_stage_interactions.py stage_crosslinks', **common)

    rpath = folder / 'rbdmap_castello2016_table_s1.xlsx'
    wb = openpyxl.load_workbook(rpath, read_only=True)
    rows = wb['RBDpep'].iter_rows(values_only=True)
    head = next(rows)
    ix = {str(h).strip(): k for k, h in enumerate(head)}
    frag_col = next(k for h, k in ix.items() if h.startswith('LysC/ArgC'))
    rbd = [[str(r[ix['ProtID']] or ''), str(r[ix['Symbol']] or ''), r[ix['fragmentStart']], r[ix['fragmentStop']], str(r[frag_col] or '')]
           for r in rows]
    d = dates_of([rpath])
    w.table('rbdmap_castello2016_rna_bound_fragments.tsv', ['protein_id', 'gene', 'fragment_start', 'fragment_end', 'fragment_sequence'], rbd,
            content='RBDmap (HeLa) high-confidence RNA-bound fragments (sheet RBDpep): the LysC/ArgC fragment next to the crosslinked peptide, '
                    'as Table S1 gives it', source=x37.RBDMAP_S1, version=f'Table S1, downloaded {d}',
            citation='Castello et al., Mol Cell 63:696 (2016), doi:10.1016/j.molcel.2016.06.029',
            licence='CC BY 4.0 (open access article and its supplement; PMC5003815)',
            how='scripts/37a_stage_interactions.py stage_crosslinks (one row per table row)', derived=['crosslinks/rbdmap_castello2016_table_s1.xlsx'])

    loc, pep, interactome = [], [], []
    p2, p1 = folder / 'trendel2025_mmc2.xlsx', folder / 'trendel2025_mmc1.xlsx'
    wb = openpyxl.load_workbook(p2, read_only=True)
    rows = wb['4ST-xl-peptides localization'].iter_rows(values_only=True)
    ix = {h: k for k, h in enumerate(next(rows))}
    for r in rows:
        loc.append([x37.base_acc(str(r[ix['Uniprot ID']])), str(r[ix['Uniprot ID']]), r[ix['Protein position start']], str(r[ix['Sequence with localization']])])
    rows = wb['4ST-xl-peptides combined'].iter_rows(values_only=True)
    ix = {h: k for k, h in enumerate(next(rows))}
    for r in rows:
        pep.append([x37.base_acc(str(r[ix['Uniprot ID']])), str(r[ix['Uniprot ID']]), r[ix['Protein position start']], r[ix['Protein position end']]])
    wb = openpyxl.load_workbook(p1, read_only=True)
    rows = wb['DNA interactome & subproteomes'].iter_rows(values_only=True)
    ix = {h: k for k, h in enumerate(next(rows))}
    for r in rows:
        if r[ix['MCF7 DNA interactome']] == 'DNA interactome':
            interactome.append([x37.base_acc(str(r[ix['Uniprot ID']])), str(r[ix['Uniprot ID']])])
    d = dates_of([p1, p2])
    common = dict(source=x37.TRENDEL_MMC.format(k='{1,2}'), version=f'Tables S1-S2, downloaded {d}',
                  citation='Trendel et al., Cell 188:4424 (2025), doi:10.1016/j.cell.2025.04.037',
                  licence='CC BY-NC 4.0 (open access article and its supplement)')
    w.table('trendel2025_dna_crosslink_localisation.tsv', ['uniprot_accession', 'uniprot_id', 'peptide_start', 'sequence_with_localization'], loc,
            content='DNA-crosslinked peptides (MCF7, sheet 4ST-xl-peptides localization): lower-case letters are the residues the crosslink '
                    'may sit on; one lower-case letter localises it to that residue', how='scripts/37a_stage_interactions.py stage_crosslinks',
            derived=['crosslinks/trendel2025_mmc2.xlsx'], **common)
    w.table('trendel2025_dna_crosslinked_peptides.tsv', ['uniprot_accession', 'uniprot_id', 'peptide_start', 'peptide_end'], pep,
            content='DNA-crosslinked peptides (MCF7, sheet 4ST-xl-peptides combined): protein ranges', how='scripts/37a_stage_interactions.py stage_crosslinks',
            derived=['crosslinks/trendel2025_mmc2.xlsx'], **common)
    w.table('trendel2025_mcf7_dna_interactome.tsv', ['uniprot_accession', 'uniprot_id'], interactome,
            content='Proteins of the MCF7 DNA interactome (sheet DNA interactome & subproteomes)', how='scripts/37a_stage_interactions.py stage_crosslinks',
            derived=['crosslinks/trendel2025_mmc1.xlsx'], **common)


# ----------------------------------------------------------------------------- CD-CODE
def stage_cdcode(w: Writer, c: Path):
    folder = c / 'cdcode'
    version = json.loads((folder / 'release.json').read_text()).get('latest_release_version', '')
    rows, files = [], []
    for p in sorted(folder.glob('*.json')):
        if p.name == 'release.json':
            continue
        files.append(p)
        try:
            data = json.loads(p.read_text()).get('data', [])
        except json.JSONDecodeError:
            continue
        for r in data:
            acc = x37.base_acc((r.get('protein') or {}).get('uniprot_id') or '')
            if not acc:
                continue
            roles, evid, pmids = [], [], []
            for a in (r.get('new_annotations') or []) + (r.get('old_annotations') or []):
                an = a.get('annotations') or {}
                t = an.get('functional_type')
                for x in (t if isinstance(t, list) else [t] if t else []):
                    if x not in roles:
                        roles.append(x)
                for e in an.get('exp_evidence') or []:
                    e = e if str(e).strip() else 'unspecified'   # CD-CODE sometimes lists an empty evidence value
                    if e not in evid:
                        evid.append(e)
                for pm in ([a.get('publication')] if a.get('publication') else []) + list(a.get('publications') or []):
                    pm = str(pm)
                    if pm.isdigit() and pm not in pmids:
                        pmids.append(pm)
            rows.append([p.stem, acc, r.get('confidence_score') or 0, [text(x) for x in roles], [text(x) for x in evid], pmids])
    d = dates_of(files)
    w.table(f'cdcode_{version}_condensate_members_{d[:10]}.tsv',
            ['condensate_uid', 'uniprot_accession', 'confidence_score', 'roles', 'experimental_evidence', 'pmids'], rows,
            content=f'CD-CODE {version} members of the condensates the pipeline uses: CD-CODE confidence score (0-5), functional roles '
                    '(member, driver, client, regulator ...), kinds of experimental evidence ("unspecified" where CD-CODE gives an '
                    'empty value) and the PubMed IDs of the annotations',
            source=f'{x37.CDCODE_API}/condensates/<uid>/proteins', version=f'CD-CODE {version} API, answers saved {d}',
            citation='Rostam et al., Nat Methods 20:673 (2023)', licence='CC BY-NC 4.0 (stated on cd-code.org)',
            how='scripts/37a_stage_interactions.py stage_cdcode', derived=[f'cdcode/<uid>.json ({len(files)} files)', 'cdcode/release.json'])


# ----------------------------------------------------------------------------- raw cache list
def list_cache(c: Path, out: Path):
    rows = []
    for p in sorted(c.rglob('*')):
        if p.is_file():
            rows.append([str(p.relative_to(c)), p.stat().st_size, md5_file(p), file_date(p)])
    with open(out / 'raw_cache_files.tsv', 'w', encoding='utf-8', newline='\n') as fh:
        fh.write('path\tbytes\tmd5\tsaved\n')
        for r in rows:
            fh.write('\t'.join(map(str, r)) + '\n')
    print(f'raw_cache_files.tsv: {len(rows):,} files, {sum(r[1] for r in rows) / 1e6:.0f} MB')


def stage_measured_positions(w: Writer, c: Path):
    """Which positions the top-up measured in each predicted model (the contacts table lists only those near the partner)."""
    rows, files = [], []
    for source, folder in (('ProtVar', 'protvar_measure'), ('AlphaFold DB', 'afdb_measure')):
        for p in sorted((c / folder).glob('*.json')):
            files.append(p)
            for ch, (acc, positions) in sorted(json.loads(p.read_text()).get('positions', {}).items()):
                rows.append([source, p.stem, ch, acc, sorted(positions)])
    d = dates_of(files)
    w.table(f'predicted_model_measured_positions_{d[:10]}.tsv', ['source', 'model', 'chain', 'uniprot_accession', 'positions'], rows,
            content='The methylarginine positions measured in each predicted model by the top-up (predicted_model_arginine_contacts '
                    'lists only those within 5 Å of the partner), so scripts/37 --tables does not measure them again',
            source=f'{x37.PROTVAR_MODEL}; {x37.AFDB_PREDICTION}', version=f'models read {d}', citation='', licence='CC BY 4.0',
            how='scripts/37a_stage_interactions.py stage_measured_positions', derived=[f'protvar_measure/ and afdb_measure/ ({len(files):,} files)'])


def stage_topup(c: Path, out: Path, workers: int, manifest: Path, measured: str):
    """What scripts/37 --tables fetched into its top-up cache (proteins and sites the frozen tables do not cover), as dated
    tables next to the frozen ones; scripts/37 --tables reads every dated table of a kind. Adds their rows to the
    inputs MANIFEST and lists the fetched files in raw_topup_files_DATE.tsv."""
    w = Writer(out, c, keep_existing=True)
    if (c / 'pdbe_interfaces').is_dir():
        stage_pdbe_interfaces(w, c)
    if (c / 'contacts_seqres').is_dir():
        stage_contacts(w, c)
    if (c / 'environments').is_dir():
        stage_environments(w, c, workers)
    if any((c / 'pdbe_molecules').glob('*.json')):
        stage_nucleic_chains(w, c)
    if any((c / 'na_regions').glob('*.json')):
        stage_na_regions(w, c)
    if any((c / 'na_arginines').glob('*.json')):
        stage_na_arginines(w, c)
    # scripts/37 makes these folders before it knows whether a model needs measuring: stage only what holds files
    if any(p.is_file() for x in ('protvar_measure', 'afdb_measure', 'afdb_complex') for p in (c / x).glob('*')):
        stage_model_contacts(w, c, served=out / 'protvar_served_models_af2complexes_v1.tsv', measured=measured)
        stage_measured_positions(w, c)
    files = [p for p in sorted(c.rglob('*')) if p.is_file()]
    raw = out / w.free(f'raw_topup_files_{dates_of(files)[:10]}.tsv')
    with open(raw, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write('path\tbytes\tmd5\tsaved\n')
        for p in files:
            fh.write(f'{p.relative_to(c)}\t{p.stat().st_size}\t{md5_file(p)}\t{file_date(p)}\n')
    w._add(raw, len(files), f'Every file the top-up fetched ({c.name}/), with its size, MD5 and date', 'the sources above', dates_of(files),
           '', '', 'scripts/37a_stage_interactions.py --topup', [], 'git')
    # the inputs MANIFEST has its own column order
    m = manifest.read_text(encoding='utf-8').splitlines()
    head = m[0].split('\t')
    names = {r[0] for r in w.manifest}
    keep = [line for line in m[1:] if line.split('\t')[0] not in names]
    for r in w.manifest:
        rec = dict(zip(MANIFEST_COLUMNS, r))
        rec['how_made_or_checked'] = rec.pop('how_made')
        keep.append('\t'.join(cell(rec.get(h, '')) for h in head))
    manifest.write_text('\n'.join([m[0]] + keep) + '\n', encoding='utf-8')
    print(f'{len(w.manifest)} rows in {manifest}')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--cache', default='PTM_data/interactions_cache')
    ap.add_argument('--out', default='inputs/interactions')
    ap.add_argument('--topup', default='', help='stage only this top-up cache of scripts/37 --tables (e.g. PTM_data/interactions_topup)')
    ap.add_argument('--manifest', default='inputs/MANIFEST.tsv', help='with --topup: the inputs MANIFEST to add the new tables to')
    ap.add_argument('--measured', default='the v1.1.0 sites that the published union does not have',
                    help='with --topup: whose positions the top-up measured, for the MANIFEST')
    ap.add_argument('--workers', type=int, default=4, help='processes for the interface-geometry measures')
    ap.add_argument('--uniprot-release', default='2026_03',
                    help='UniProt release of the cached UniProt downloads, recorded in the MANIFEST (the cache does not store it)')
    ap.add_argument('--reference-fasta', default='inputs/reference/uniprot_human_reviewed_2026_03.fasta',
                    help='reference FASTA whose entries are marked in_reference_proteome (empty: leave the column empty)')
    args = ap.parse_args()
    if args.topup:
        stage_topup(Path(args.topup), Path(args.out), args.workers, Path(args.manifest), args.measured)
        return
    c, out = Path(args.cache), Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    w = Writer(out, c)
    reference = set(x37.read_fasta(Path(args.reference_fasta))) if args.reference_fasta else set()
    seqs, uniprot_rows = stage_uniprot(w, c, args.uniprot_release, reference)
    reviewed = set(seqs)
    stage_go(w, c)
    stage_string(w, c, uniprot_rows)
    stage_complexes(w, c)
    stage_pdbe(w, c, reviewed)
    stage_modelled(w, c)
    stage_contacts(w, c)
    stage_environments(w, c, args.workers)
    stage_protvar(w, c)
    stage_model_contacts(w, c)
    stage_intact(w, c, seqs)
    stage_crosslinks(w, c)
    stage_cdcode(w, c)
    w.write_manifest()
    # Tables whose licence does not yet allow publishing are written but kept out of Git until it is settled.
    held = sorted(Path(r[0]).name for r in w.manifest if r[-1] == 'held_pending_licence')
    if held:
        (out / '.gitignore').write_text('# held until their licence is settled (MANIFEST.tsv redistribution = held_pending_licence)\n'
                                        + ''.join(f'{n}\n' for n in held))
    elif (out / '.gitignore').exists():
        (out / '.gitignore').unlink()
    list_cache(c, out)


if __name__ == '__main__':
    main()
