#!/usr/bin/env python3
"""Stage the PRMT5 inhibitor (EPZ015666) arginine-methylation response of Lu et al. 2024 on the canonical human proteome.

    python scripts/51_stage_lu2024_prmt5.py --xlsx pr3c00724_si_002.xlsx --fasta inputs/reference/uniprot_human_reviewed_2026_03.fasta \
        --outdir inputs/layers

Lu, Ye, Zhang, Olsen, Yuan & Mao, "ETD-Based Proteomic Profiling Improves Arginine Methylation Identification and
Reveals Novel PRMT5 Substrates", J Proteome Res 23:1014-1027 (2024), doi:10.1021/acs.jproteome.3c00724. Supporting
Information 2 (pr3c00724_si_002.xlsx, doi:10.1021/acs.jproteome.3c00724.s002, CC BY-NC 4.0 on ACS's figshare portal)
Table 1.1: the 391 methyl peptide groups quantified by SILAC in HeLa cells with and without the PRMT5 inhibitor
EPZ015666 (ulilysin digest, PTMScan sdme-RG enrichment, Proteome Discoverer 2.4 with Minora; PRIDE PXD037425), each
with its marks and ptmRS probabilities, its master protein and span, heavy and light abundances of three replicates
(F4-F6), LH1-LH3 = log2 light/heavy per replicate, LH = their median, a p value and padj.

Which channel was treated is not stated in any text that could be read (the article's methods are behind a bot check).
Table 1.5 calls the Table 1.4 sites 'downregulated' and every one but one has a peptide at LH >= 1 (light above heavy),
so the heavy channel is read as the EPZ015666-treated one: log2_inhibitor_control = -LH. The threshold the paper uses
is not stated either; the rule 'a peptide at nominal p < 0.05 and LH >= 1' (twofold down) reproduces Table 1.4 for 255
of its 256 sites and adds none (P19338 R691 is the one left: its peptides are LH 0.54, p 0.37 and LH -4.62, p 0.017). The
call here is that rule, as scripts/41 calls the CARM1 knockout (down: log2 <= -1, up: >= 1, both p < 0.05), and the
paper's own list is kept beside it (paper_prmt5_regulated).

Each peptide is placed on the canonical sequence of its master protein's base accession (scripts/41 place(): it must
occur there exactly once). Its methylarginines are the marks ptmRS places (any probability Proteome Discoverer shows,
>= 75 %; as Table 1.2 counts them), lower-cased in peptide; methyl groups ptmRS does not place are counted in
unplaced_methyl_groups.

Outputs (in --outdir): lu2024_prmt5_peptides.tsv, lu2024_prmt5_sites.tsv, lu2024_prmt5_summary.json.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
from collections import Counter, namedtuple
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))


def load_script(name: str, file: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / file)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


s41 = load_script('stage_carm1_ko_sites', '41_stage_carm1_ko_sites.py')
rl = load_script('stage_study_sites', '00_stage_study_sites.py')   # the study's reader: its Proteome Discoverer mark parsers

STUDY, DOI = 'Lu 2024', '10.1021/acs.jproteome.3c00724'
CELL = 'HeLa'
RULE = (f'peptide log2 EPZ015666/control (= -LH) <= -{s41.LOG2_CUTOFF:g} (down) or >= {s41.LOG2_CUTOFF:g} (up) with nominal '
        f'p < {s41.P_CUTOFF:g} (scripts/41 call); reproduces 255 of the 256 sites of Table 1.4 and adds none')
Row = namedtuple('Row', 'source_UniProtAC peptide reported_start')


def sitelist(wb, sheet: str) -> set[tuple[str, int]]:
    return set(rl.lu2024_sitelist(wb, sheet))


def genes(wb) -> dict[str, str]:
    """Accession -> gene symbol, as Tables 1.5 and 1.6 give them (the proteins with a PRMT5-regulated site)."""
    out = {}
    for r in list(wb['Table1.5'].iter_rows(values_only=True))[1:]:
        if r and rl.text(r[0]) and rl.text(r[1]):
            out[rl.text(r[0])] = rl.text(r[1])
    for r in list(wb['Table1.6'].iter_rows(values_only=True))[2:]:
        if r and rl.text(r[1]) and rl.text(r[0]):
            out.setdefault(rl.text(r[1]), rl.text(r[0]))
    return {a: g for a, g in out.items() if g != a}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--xlsx', required=True, help='Lu 2024 Supporting Information 2 (pr3c00724_si_002.xlsx)')
    ap.add_argument('--fasta', required=True, help='the reviewed canonical FASTA')
    ap.add_argument('--outdir', required=True)
    args = ap.parse_args()
    xlsx = Path(args.xlsx)
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    canonical = s41.parse_fasta(args.fasta)
    rows = list(wb['Table1.1'].iter_rows(values_only=True))
    H = {rl.text(h): k for k, h in enumerate(rows[1]) if h is not None}
    data = [r for r in rows[2:] if r and rl.text(r[0])]
    gene_of = genes(wb)
    table_1_2, table_1_4 = sitelist(wb, 'Table1.2'), sitelist(wb, 'Table1.4')
    floor = min(float(r[H[c]]) for r in data for c in H if c.startswith('Abundance F') and r[H[c]] is not None)
    peps, reported_to_canonical, stats = [], {}, Counter()
    for r in data:
        acc = rl.text(r[H['Master Protein Accessions']])
        plain = re.sub(r'^\[.*?\]\.|\.\[.*?\]$', '', rl.text(r[H['Annotated Sequence']])).upper()
        placed, unplaced = rl.lu2024_marks(r[H['Modifications']])
        kinds = Counter((res, off) for res, off, k, _ in placed if k in ('Methyl', 'Dimethyl', 'Trimethyl'))
        marks = {off: rl.DEGREE[k] for res, off, k, _ in placed
                 if res == 'R' and k in rl.DEGREE and kinds[(res, off)] == 1}
        starts = rl.lu2024_starts(r[H['Positions in Master Proteins']], acc)
        status, start = s41.place(Row(acc, plain, starts[0] if len(starts) == 1 else ''), canonical)
        lh, p, padj = r[H['LH']], r[H['pvalue']], r[H['padj']]
        log2 = None if lh is None else -float(lh)
        call = s41.call(log2, float(p)) if log2 is not None and p is not None else 'not_quantified'
        positions = sorted(marks)
        peptide = ''.join(ch.lower() if k + 1 in marks else ch for k, ch in enumerate(plain))
        abund = [r[H[c]] for c in H if c.startswith('Abundance F')]
        peps.append({
            'cell_line': CELL, 'peptide': peptide, 'reported_start': starts[0] if len(starts) == 1 else '',
            'source_UniProtAC': acc, 'gene': gene_of.get(acc, ''), 'log2_inhibitor_control': log2, 'p': p,
            'placement': status, 'canonical_start': start, 'canonical_UniProtAC': acc.split('-')[0], 'call': call,
            'methyl_positions': ';'.join(str(start + off - 1) for off in positions) if start else '',
            # beyond the columns of carm1_ko_peptides.tsv
            'methyl_degrees': ';'.join(marks[off] for off in positions),
            'unplaced_methyl_groups': sum(unplaced.values()), 'padj': padj, 'log2_light_heavy': lh,
            'log2_light_heavy_replicates': ';'.join(f'{float(r[H[c]]):.6g}' for c in ('LH1', 'LH2', 'LH3')),
            'abundances_at_table_minimum': sum(1 for v in abund if v is not None and float(v) == floor),
        })
        if start and len(starts) == 1:
            for off in positions:
                reported_to_canonical[(acc, starts[0] + off - 1)] = (acc.split('-')[0], start + off - 1)
    peptides = pd.DataFrame(peps)
    peptides['canonical_start'] = peptides['canonical_start'].astype('Int64')

    site_rows = []
    for row in peptides.itertuples(index=False):
        if pd.isna(row.canonical_start) or row.call == 'not_quantified':
            continue
        for pos, deg in zip(row.methyl_positions.split(';'), row.methyl_degrees.split(';')):
            if pos:
                site_rows.append({'canonical_UniProtAC': row.canonical_UniProtAC, 'position': int(pos), 'degree': deg,
                                  'log2': row.log2_inhibitor_control, 'p': float(row.p), 'padj': float(row.padj),
                                  'call': row.call})
    per_site = pd.DataFrame(site_rows)
    paper_4 = {reported_to_canonical[s] for s in table_1_4 if s in reported_to_canonical}
    paper_2 = {reported_to_canonical[s] for s in table_1_2 if s in reported_to_canonical}
    off_canonical = sorted(f'{a}:R{p}' for a, p in table_1_2 if (a, p) not in reported_to_canonical)
    sites = []
    for (acc, pos), grp in per_site.groupby(['canonical_UniProtAC', 'position'], sort=True):
        calls = Counter(grp['call'])
        site_call = 'mixed' if calls['down'] and calls['up'] else 'down' if calls['down'] else 'up' if calls['up'] else 'unchanged'
        sites.append({'canonical_UniProtAC': acc, 'position': pos, 'residue': canonical[acc][pos - 1], 'cell_line': CELL,
                      'median_log2_inhibitor_control': round(float(grp['log2'].median()), 3), 'min_p': float(grp['p'].min()),
                      'peptides': len(grp), 'peptides_down': calls['down'], 'peptides_up': calls['up'], 'call': site_call,
                      'degrees': ';'.join(sorted(set(grp['degree']))), 'min_padj': float(grp['padj'].min()),
                      'paper_quantified_rme2': (acc, pos) in paper_2, 'paper_prmt5_regulated': (acc, pos) in paper_4})
    sites = pd.DataFrame(sites)
    not_r = sites[sites['residue'].ne('R')]
    if len(not_r):
        raise SystemExit(f'{len(not_r)} placed methylation sites are not arginines on the canonical sequence:\n{not_r}')
    rme2 = sites[sites['degrees'].str.contains('Rme2')]
    down = set(map(tuple, rme2.loc[rme2['call'].isin(['down', 'mixed']), ['canonical_UniProtAC', 'position']].values))
    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    s41.save_table(peptides, outdir / 'lu2024_prmt5_peptides.tsv')
    s41.save_table(sites, outdir / 'lu2024_prmt5_sites.tsv')
    summary = {
        'study': STUDY, 'doi': DOI, 'rule': RULE,
        'treated_channel': 'heavy (inferred: Table 1.5 calls the Table 1.4 sites downregulated, and they have light > heavy)',
        'input': xlsx.name, 'input_md5': hashlib.md5(xlsx.read_bytes()).hexdigest(),
        'input_sha256': hashlib.sha256(xlsx.read_bytes()).hexdigest(), 'canonical_fasta': Path(args.fasta).name,
        'peptides': len(peptides), 'placement': dict(Counter(peptides['placement'])),
        'peptide_calls': dict(Counter(peptides['call'])),
        'sites': dict(Counter(sites['call'])), 'unique_sites': len(sites),
        'sites_rme2': len(rme2), 'paper_table_1_2_sites': len(table_1_2), 'paper_table_1_4_sites': len(table_1_4),
        'paper_table_1_2_not_placed_on_canonical': off_canonical,
        'paper_table_1_4_on_canonical': len(paper_4),
        'paper_table_1_2_sites_without_a_quantified_peptide': sorted(
            f'{a}:R{p}' for a, p in paper_2 - set(map(tuple, sites[['canonical_UniProtAC', 'position']].values))),
        'rme2_sites_down_or_mixed_by_rule': len(down),
        'rule_vs_table_1_4': {'both': len(down & paper_4), 'rule_only': sorted(f'{a}:R{p}' for a, p in down - paper_4),
                              'table_1_4_only': sorted(f'{a}:R{p}' for a, p in paper_4 - down)},
        'abundance_minimum_of_table': floor,
        'peptides_with_an_abundance_at_the_minimum': int((peptides['abundances_at_table_minimum'] > 0).sum()),
    }
    (outdir / 'lu2024_prmt5_summary.json').write_text(json.dumps(summary, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
