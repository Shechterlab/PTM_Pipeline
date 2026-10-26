#!/usr/bin/env python3
"""Stage the PRMT7, PRMT4 (CARM1) and PRMT5 knockdown and SGC3027 arginine-methylation responses (Li et al. 2021) on the
canonical human proteome.

Li, He, Yang, Hu, Lin, Ran, Peng, Xie, Huang, Gao, Huang, Zhu, Ye & Liu, "Profiling PRMT methylome reveals roles of hnRNPA1
arginine methylation in RNA splicing and cell growth", Nat Commun 12:1946 (2021), doi:10.1038/s41467-021-21963-1. The
article is CC BY 4.0; its Supplementary Data are taken as covered by that licence (the user's decision for Nature
Communications supplements), so the staged tables are redistributed in Git. Raw MS files: PRIDE PXD022424, not used.

HEK293 cells, SILAC (Methods, "Cell labeling by SILAC"): "Light and heavy labeled cells were transfected with control siRNA
and siRNA specifically targeting PRMT4, 5 or 7, respectively", so heavy is the knockdown (Results: "SILAC labeled wild-type
(light) or PRMT7-knockdown (heavy) cells (Fig. 1d)"). Peptides were enriched with the PTMScan monomethylarginine antibodies
(10 fractions), then the supernatant with ADMA or SDMA antibodies (5 fractions), and searched in Proteome Discoverer 2.2
with methylation of R and K. Four Supplementary Data workbooks are read, each sheet once:

  Supplementary Data 1 (MOESM3)  siPRMT7: 'repeat 1' and 'repeat 2' (two Proteome Discoverer exports of two SILAC
                                 experiments) and the paper's lists 'PRMT7-regulated me1 (FC>2)' and '(abolished)'
  Supplementary Data 3 (MOESM5)  SGC3027 (PRMT7 inhibitor): 'Raw data' and the lists 'SGC3027-regulated me1 (FC>2)' and
                                 'SGC3027regulated me1(abolished)' (their headers say siPRMT7/siCTL, copied)
  Supplementary Data 5 (MOESM7)  siPRMT4: 'raw data' and the lists me1 and me2, (FC>2) and (abolished)
  Supplementary Data 6 (MOESM8)  siPRMT5: the lists me1 and me2, (FC>2) and (abolished); its 'raw data' sheet is read
                                 only to check it (below)

Each export is a Proteome Discoverer hierarchy (protein, peptide group, and for SD1 and SD3 PSM rows). A peptide group
gives the peptide, its modifications as PD placed them ('1xMethyl [R11]', or '1xMethyl [R]' / '[R/K]' when PD did not
localise the mark), its master proteins, # PSMs, 'Abundance Ratio: (Heavy) / (Light)' (PD floors it at 0.01, when the
heavy channel has no signal, and caps it at 100) and, in SD1 and SD3 only, 'Abundance Ratio Adj. P-Value'. The paper
does not say how PD computed that adjusted p; each export compares one SILAC pair of cultures (SD1 repeat 1 enters its 15
raw files as 15 sample pairs), so it is not a test over biological replicates. It is kept as given.

The lists give, per site, the accession, gene, unmarked peptide, the position in the protein and the peptide, # PSMs,
the methylation ratio, a protein-level ratio from a separate proteome and their quotient ('normalized ratio'). Every list
row is a peptide group of the export with the same accession, sequence, ratio and # PSMs (SD1 repeat 1, SD3, SD5; SD1
repeat 2 holds 19 of the 245 listed PRMT7 peptides, so the PRMT7 lists come from repeat 1). The paper calls on the
methylation ratio, not the normalised one: "297 MMA sites ... had at least a two-fold reduction ... when methylation was
normalized to protein levels, only the change in 25 methylation sites (from 20 proteins) was due to the change in protein
levels" (the list holds those 25 rows with normalised ratio > 0.5). 'Abolished' ("completely abolished") rows are the
listed rows at the 0.01 floor. The sheet names say FC>2, the text "at least two-fold"; of the two methyl-arginine peptide
groups at exactly 0.5, the one whose site no other group lists (SD5, Q8TD08 R449) is not listed, so the paper's rule is read
as H/L < 0.5. SD6's 'raw data' sheet is not the export the PRMT5 lists come from (10,282 master proteins; 34 of the 337
listed peptides among its 422 methyl-arginine peptide groups, and 26 of the 843 list rows' ratios: see the summary), so
PRMT5 has the listed sites only, and no quantified-but-unlisted site.

SGC3027: the Methods say "light and heavy labeled cells were treated with or without SGC3027, respectively", which read
literally makes the heavy channel the untreated one; the paper reads the H/L values of its SGC3027 lists as decreases
"upon SGC3027 treatment", and the SGC3027 H/L ratios agree in direction with the siPRMT7 ones (the summary counts the
peptide groups changed at least twofold in both). The stage follows the paper: log2 treated/control = log2 H/L, heavy =
SGC3027, flagged in the summary.

Comparisons: PRMT7_si (SD1 repeat 1, the paper's), PRMT7_si_rep2 (SD1 repeat 2: the paper makes no call from it, so its
rows carry the uniform call only), PRMT7_SGC3027, PRMT4_si and PRMT5_si (lists only).

Placement. Each quantified methyl-arginine peptide group (H/L given) is placed on the canonical sequence of its first master
protein, or of the master the lists name when they put a shared peptide on another only: it must occur there exactly once
(position_verified when that is PD's start on the master, else placed_by_sequence). The lists name both entries of a few
shared peptides (TNPO1/TNPO2, RBMX/RBMXL1, EP300/CREBBP, SNRPB/SNRPN, POM121C/POM121, CSTF2/CSTF2T); those are staged once,
on the first master, and the summary lists the other entry's sites. Its sites are the arginines PD localised a methyl mark to, and those the paper's lists give for it
(the lists name every arginine of a peptide whose mark PD did not localise, as scripts/00_stage_study_sites.py reads them);
positions are kept only where the canonical sequence has an arginine (common.residue_is). A list row not traced to a group
(all of PRMT5) is a peptide of its own, placed the same way on its accession; its sites are the listed positions.

Calls. call_source (the paper's rule, per peptide): decreased at H/L < 0.5, else no_significant_change, for the arms the
lists cover and the degrees they list (PRMT7 arms: Rme1; PRMT4, PRMT5: Rme1 and Rme2); abolished: H/L at the 0.01 floor.
At a site, call_source is the paper's call: decreased when a list (FC>2 or abolished, me1 or me2) names the site, else
no_significant_change when a peptide in scope carries it, else empty; call_source_rule_peptides is the paper's rule over
its peptides (the two differ where the paper listed one peptide group per site). call_uniform (where the export gives an
adjusted p: PRMT7_si, PRMT7_si_rep2, PRMT7_SGC3027) is the release's rule: a peptide decreases at log2 <= -1 and adjusted
p < 0.05, increases at log2 >= 1 and adjusted p < 0.05; a site takes the call of its peptides by scripts/peptide_calls.py,
and the call of its single-site peptides alone.

Outputs (in --outdir)
  li2021_peptides.tsv     every quantified methyl-arginine peptide group of the exports, and every list peptide not in
                          one, with placement, methyl positions and degree, log2 treated/control, adjusted p, the lists
                          naming it, abolished, call_source and call_uniform
  li2021_sites.tsv        one row per canonical site and comparison: the paper's call, abolished, the listed ratios, and
                          the calls over the peptides
  li2021_summary.json     input MD5s, orientation, rules, counts per comparison against the paper's, rows left out and why
  manifest_rows.tsv       the three tables' rows for inputs/MANIFEST.tsv (as layers/...)
  raw_files_rows.tsv      the inputs' rows for inputs/layers/raw_files.tsv (pointing at their MANIFEST rows)
Nothing outside --outdir is written: --manifest is read only, to check the inputs' MD5s and take their URL and date.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import peptide_calls as pc
from common import md5_file, parse_fasta, residue_is, save_table

STUDY = 'Li 2021'
DOI = '10.1038/s41467-021-21963-1'
CITATION = 'Li WJ et al., Nat Commun 12:1946 (2021), doi:10.1038/s41467-021-21963-1'
LICENCE = 'CC BY 4.0 (the article; its Supplementary Data are taken as covered by the article\'s licence)'
CELL_LINE = 'HEK293'
SOURCE_CUTOFF = 0.5          # the paper's rule: H/L < 0.5 (sheets 'FC>2'; text 'at least two-fold')
FLOOR, CEILING = 0.01, 100.0  # Proteome Discoverer's ratio limits; the floor is the paper's 'abolished'
Q_CUTOFF, LOG2_CUTOFF = 0.05, 1.0
SOURCE_RULE = (f'peptide H/L < {SOURCE_CUTOFF:g} (more than twofold lower with the perturbation): decreased; abolished at the '
               f'Proteome Discoverer floor H/L = {FLOOR:g}. Site: decreased when the paper\'s lists (FC>2 or abolished, me1 or me2) '
               'name it, no_significant_change when a quantified peptide in the lists\' scope carries it')
UNIFORM_RULE = (f'peptide log2 treated/control <= -{LOG2_CUTOFF:g} (decreased) or >= {LOG2_CUTOFF:g} (increased) with the export\'s '
                f'adjusted p < {Q_CUTOFF:g}; sites by scripts/peptide_calls.site_response (strict)')
PAPER_RULE = ('"297 MMA sites from 174 proteins had at least a two-fold reduction in mono-methylation signals" and "methylation at '
              '176 MMA sites from 108 proteins was completely abolished when PRMT7 was knocked down" (Results); sheet names '
              '"(FC>2)" and "(abolished)"; no statistical test is stated for the ratios')
ORIENTATION = {
    'knockdown': ('heavy = siRNA against the PRMT, light = control siRNA: "Light and heavy labeled cells were transfected with control '
                  'siRNA and siRNA specifically targeting PRMT4, 5 or 7, respectively" (Methods, Cell labeling by SILAC); "SILAC labeled '
                  'wild-type (light) or PRMT7-knockdown (heavy) cells (Fig. 1d)" (Results). The lists\' ratio (siPRMTx/siCTL) is the '
                  'export\'s Abundance Ratio (Heavy)/(Light) of the same peptide group'),
    'inhibitor': ('heavy = SGC3027, light = untreated, as the paper reads its SGC3027 lists ("Upon SGC3027 treatment, 503 MMA sites ... had '
                  'at least two-fold reduction"; the listed ratios are the export\'s H/L). The Methods sentence "light and heavy labeled '
                  'cells were treated with or without SGC3027, respectively" read literally says the reverse; the list headers say '
                  'siPRMT7/siCTL (copied from SD1). See sgc3027_orientation_check'),
}
NORMALISATION = ('None applied here. The export\'s Abundance Ratio (Heavy)/(Light) is Proteome Discoverer 2.2\'s (its normalisation '
                 'settings are not stated). The lists add a protein-level ratio from a separate proteome and the "normalized ratio" = '
                 'methylation ratio / protein ratio; the paper calls on the methylation ratio and uses the normalised one only to say '
                 'how many changes protein levels explain (25 PRMT7 sites)')
REPLICATES = ('Not stated in the paper for the SILAC methylomes. SD1 holds two exports, "repeat 1" (raw files siPRMT7-Methylation-SILAC-1 '
              'to 10 and -Dimethylation-SILAC-1 to 5, as 15 light/heavy sample pairs) and "repeat 2" (AD293-sIPRMT7-SILAC-..., one '
              'sample pair); the PRMT7 lists come from repeat 1 only. SD3 and SD5 hold one export each (one light/heavy sample pair)')
PAPER = {
    'PRMT7_si': {'sites_detected': 1031, 'proteins_detected': 513, 'sites_quantified': 939, 'sites_decreased': 297,
                 'proteins_decreased': 174, 'sites_abolished': 176, 'proteins_abolished': 108,
                 'explained_by_protein_level': '25 sites from 20 proteins',
                 'where': 'Results, "Global profiling of PRMT7 substrates in HEK293 cells"; Fig. 1e (MMA sites)'},
    'PRMT7_SGC3027': {'sites_detected': 1126, 'proteins_detected': 524, 'sites_quantified': 1027, 'sites_decreased': 503,
                      'proteins_decreased': 274, 'where': 'Results, "Global profiling of PRMT7 substrates in HEK293 cells"; '
                                                         'Supplementary Fig. 1b (MMA sites; no abolished count in the text)'},
    'PRMT4_si': {'sites_detected': 1504, 'mma_sites_detected': 1036, 'adma_sites_detected': 588, 'proteins_detected': 665,
                 'sites_decreased': 660, 'proteins_decreased': 301, 'sites_abolished': 592, 'proteins_abolished': 276,
                 'where': 'Results, "Features of arginine methylation mediated by PRMT members"; Fig. 3a (mono- and dimethyl sites '
                          'after removing duplicates; the quantified count is in Fig. 3a only)'},
    'PRMT5_si': {'sites_detected': 1216, 'mma_sites_detected': 1131, 'sdma_sites_detected': 366, 'proteins_detected': 585,
                 'sites_decreased': 429, 'proteins_decreased': 244, 'sites_abolished': 238, 'proteins_abolished': 146,
                 'where': 'Results, "Features of arginine methylation mediated by PRMT members"; Fig. 3a'},
}

SD = {'sd1': ('41467_2021_21963_MOESM3_ESM.xlsx', 'Supplementary Data 1'), 'sd3': ('41467_2021_21963_MOESM5_ESM.xlsx', 'Supplementary Data 3'),
      'sd5': ('41467_2021_21963_MOESM7_ESM.xlsx', 'Supplementary Data 5'), 'sd6': ('41467_2021_21963_MOESM8_ESM.xlsx', 'Supplementary Data 6')}


@dataclass
class Arm:
    comparison: str
    sd: str                         # which input (SD keys)
    raw_sheet: str                  # the Proteome Discoverer export the lists come from; '' when the files hold none
    lists: list = field(default_factory=list)   # (sheet, degree, kind)
    scope: tuple = ()               # degrees the paper calls in this arm ('' for no call)
    enzyme: str = ''
    perturbation: str = ''
    kind: str = 'knockdown'
    note: str = ''


ARMS = [
    Arm('PRMT7_si', 'sd1', 'repeat 1', [('PRMT7-regulated me1 (FC>2)', 'Rme1', 'fc2'), ('PRMT7-regulated me1 (abolished)', 'Rme1', 'abolished')],
        ('Rme1',), 'PRMT7', 'siPRMT7 / siCTL', note='SD1 repeat 1, the experiment the paper\'s PRMT7 lists come from'),
    Arm('PRMT7_si_rep2', 'sd1', 'repeat 2', [], (), 'PRMT7', 'siPRMT7 / siCTL',
        note='SD1 repeat 2, a second SILAC experiment: the paper reports no call from it, so call_source is empty'),
    Arm('PRMT7_SGC3027', 'sd3', 'Raw data', [('SGC3027-regulated me1 (FC>2)', 'Rme1', 'fc2'), ('SGC3027regulated me1(abolished)', 'Rme1', 'abolished')],
        ('Rme1',), 'PRMT7', 'SGC3027 / untreated', kind='inhibitor', note='SGC3027 concentration and duration not given for the profiling'),
    Arm('PRMT4_si', 'sd5', 'raw data', [('PRMT4-regulated me1 (FC>2)', 'Rme1', 'fc2'), ('PRMT4-regulated me1 (abolished)', 'Rme1', 'abolished'),
                                        ('PRMT4-regulated me2 (FC>2)', 'Rme2', 'fc2'), ('PRMT4-regulated me2 (abolished)', 'Rme2', 'abolished')],
        ('Rme1', 'Rme2'), 'CARM1 (PRMT4)', 'siPRMT4 / siCTL', note='MMA then ADMA antibody; the export gives no adjusted p'),
    Arm('PRMT5_si', 'sd6', '', [('PRMT5-regulated me1 (FC>2)', 'Rme1', 'fc2'), ('PRMT5-regulated me1 (abolished)', 'Rme1', 'abolished'),
                                ('PRMT5-regulated me2 (FC>2)', 'Rme2', 'fc2'), ('PRMT5-regulated me2 (abolished)', 'Rme2', 'abolished')],
        ('Rme1', 'Rme2'), 'PRMT5', 'siPRMT5 / siCTL',
        note='MMA then SDMA antibody; the lists only (the workbook\'s raw data sheet is not their export), so no unlisted site is quantified'),
]
SD6_CHECK_SHEET = 'raw data'
LIST_COLUMNS = ['Accession number', 'Gene name', 'Peptide sequence', 'Rme1 position in protein', 'Rme1 position in peptide', 'PSM']
RATIO_H, ADJ_H = 'Abundance Ratio: (Heavy) / (Light)', 'Abundance Ratio Adj. P-Value: (Heavy) / (Light)'
MOD = re.compile(r'(?:(\d+)x)?(Methyl|Dimethyl|Trimethyl) \[([^\]]*)\]')
DEGREE = {'Methyl': 'Rme1', 'Dimethyl': 'Rme2'}
PLACED = ('position_verified', 'placed_by_sequence')

ESM = 'https://static-content.springer.com/esm/art%3A10.1038%2Fs41467-021-21963-1/MediaObjects/'
HOW = ('scripts/50_stage_li2021_perturbations.py; peptides placed on reference/uniprot_human_reviewed_2026_03.fasta (exactly once; a site '
       'is kept only where the canonical sequence has an arginine)')
MANIFEST_ROWS = {
    'li2021_peptides.tsv': (
        'Methyl-arginine peptide groups quantified by SILAC in HEK293 cells after PRMT7, PRMT4 or PRMT5 knockdown or SGC3027 (Li et al. 2021 '
        'Supplementary Data 1, 3, 5: Proteome Discoverer exports; Supplementary Data 6: the regulated lists), each placed on the canonical '
        'sequence, with methyl positions and degree, log2 treated/control, the export\'s adjusted p where given, the lists naming it, '
        f'abolished (H/L at {FLOOR:g}), the paper\'s call (H/L < {SOURCE_CUTOFF:g}) and the uniform call (|log2| >= {LOG2_CUTOFF:g} at '
        f'adjusted p < {Q_CUTOFF:g})'),
    'li2021_sites.tsv': (
        'One row per canonical methylarginine and comparison (PRMT7_si, PRMT7_si_rep2, PRMT7_SGC3027, PRMT4_si, PRMT5_si): the paper\'s call '
        '(decreased when its FC>2 or abolished lists name the site), abolished, the listed ratios, and the calls over its peptides by the '
        'paper\'s rule and by the uniform rule (scripts/peptide_calls.py)'),
    'li2021_summary.json': (
        'Input checksums, orientation, rules, counts per comparison against the paper\'s, and the rows left out and why, of the Li 2021 '
        'perturbation staging'),
}


# ---- reading ----------------------------------------------------------------------------------------------------------

def num(v) -> float | None:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(x) else x


def text(v) -> str:
    return '' if v is None else str(v).strip()


def bare(annotated: str) -> str:
    """'[R].NPASRPPASGRGER.[A]' -> 'NPASRPPASGRGER'."""
    m = re.fullmatch(r'\[[^\]]*\]\.(.*)\.\[[^\]]*\]', annotated)
    return (m.group(1) if m else annotated).upper()


def marks(mods: str) -> tuple[dict[int, set], int, int]:
    """From a peptide group's PD modifications: {offset: degrees} of the methyl marks PD localised to an arginine, the number
    of methyl/dimethyl marks it did not localise that may sit on an arginine ('[R]', '[R/K]'), and those localised to a lysine."""
    loc, unloc, lys = defaultdict(set), 0, 0
    for n, kind, where in MOD.findall(mods):
        if kind not in DEGREE:
            continue
        toks = [t.strip() for t in where.split(';') if t.strip()]
        placed = [t for t in toks if re.fullmatch(r'[RK]\d+', t)]
        for t in placed:
            if t[0] == 'R':
                loc[int(t[1:])].add(DEGREE[kind])
            else:
                lys += 1
        rest = int(n or 1) - len(placed)
        if rest > 0 and any('R' in t for t in toks if t not in placed):
            unloc += rest
    return dict(loc), unloc, lys


def has_methyl_r(mods: str) -> bool:
    return any(kind in DEGREE and 'R' in where for _, kind, where in MOD.findall(mods))


def read_export(path: Path, sheet: str) -> tuple[list[dict], dict]:
    """The methyl-arginine peptide groups of one Proteome Discoverer export sheet (one per sequence and modifications), and
    counts of what the sheet holds."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f'{path}: no sheet {sheet!r} in {wb.sheetnames}')
    H, groups, counts = None, {}, Counter()
    all_groups, light, heavy = set(), '', ''
    for k, row in enumerate(wb[sheet].iter_rows(values_only=True), 1):
        if k == 1:
            continue
        lead = row[0] if row else None
        if lead not in (None, ''):
            counts['master_protein_rows'] += 1
            continue
        if len(row) < 2 or row[1] in (None, ''):
            counts['psm_rows'] += 1 if len(row) > 2 and row[2] not in (None, '') and row[2] != 'Checked' else 0
            continue
        if row[1] == 'Checked':
            h = {text(c): i for i, c in enumerate(row) if text(c)}
            if H is None:
                H = h
                missing = [c for c in ('Annotated Sequence', 'Modifications', 'Master Protein Accessions', 'Positions in Master Proteins',
                                       '# PSMs', RATIO_H) if c not in H]
                if missing:
                    raise SystemExit(f'{path} sheet {sheet}: peptide header lacks {missing}')
                light = next((c for c in H if c.startswith('Abundances (') and 'Light' in c and 'CV' not in c), '')
                heavy = next((c for c in H if c.startswith('Abundances (') and 'Heavy' in c and 'CV' not in c), '')
            elif h != H:
                raise SystemExit(f'{path} sheet {sheet}: two different peptide-group headers (row {k})')
            continue
        get = lambda c: row[H[c]] if c in H and H[c] < len(row) else None  # noqa: E731
        annotated, mods = text(get('Annotated Sequence')), text(get('Modifications'))
        all_groups.add((annotated, mods))
        counts['peptide_group_rows'] += 1
        if not has_methyl_r(mods):
            continue
        g = {'source_row': k, 'annotated_sequence': annotated, 'peptide': bare(annotated), 'modifications': mods,
             'master_accessions': text(get('Master Protein Accessions')), 'positions_in_masters': text(get('Positions in Master Proteins')),
             'psms': int(get('# PSMs')) if num(get('# PSMs')) is not None else None,
             'ratio_heavy_light': num(get(RATIO_H)), 'adj_p': num(get(ADJ_H)) if ADJ_H in H else None,
             'abundance_light': num(get(light)) if light else None, 'abundance_heavy': num(get(heavy)) if heavy else None,
             'quan_info': text(get('Quan Info')), 'quan_usage': text(get('Quan Usage')), 'confidence': text(get('Confidence'))}
        key = (annotated, mods)
        if key in groups:
            old = groups[key]
            if any(old[c] != g[c] for c in ('ratio_heavy_light', 'adj_p', 'psms', 'master_accessions')):
                raise SystemExit(f'{path} sheet {sheet}: peptide group {key} listed twice with different values')
            counts['repeated_under_another_master'] += 1
            continue
        groups[key] = g
    counts['peptide_groups'] = len(all_groups)
    counts['methyl_arginine_groups'] = len(groups)
    counts['adjusted_p_column'] = int(H is not None and ADJ_H in H)
    return list(groups.values()), dict(counts)


def read_lists(path: Path, arm: Arm) -> list[dict]:
    """The rows of the paper's lists of one arm."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for sheet, degree, kind in arm.lists:
        if sheet not in wb.sheetnames:
            raise SystemExit(f'{path}: no sheet {sheet!r}')
        rows = wb[sheet].iter_rows(values_only=True)
        header = [text(h) for h in next(rows)]
        missing = [c for c in LIST_COLUMNS if c not in header]
        ratio_cols = [i for i, h in enumerate(header) if h.startswith('ratio (') or h.startswith('normalized ratio (')]
        if missing or len(ratio_cols) != 3:
            raise SystemExit(f'{path} sheet {sheet}: columns {header}')
        at = {c: header.index(c) for c in LIST_COLUMNS}
        for k, r in enumerate(rows, 2):
            if not r or all(v in (None, '') for v in r):
                continue
            out.append({'sheet': sheet, 'degree': degree, 'kind': kind, 'row': k, 'accession': text(r[at['Accession number']]),
                        'gene': text(r[at['Gene name']]), 'peptide': text(r[at['Peptide sequence']]).upper(),
                        'reported_position': int(r[at['Rme1 position in protein']]), 'offset': int(r[at['Rme1 position in peptide']]),
                        'psms': int(r[at['PSM']]), 'ratio': float(r[ratio_cols[0]]), 'protein_ratio': float(r[ratio_cols[1]]),
                        'normalised_ratio': float(r[ratio_cols[2]])})
    return out


def check_sd6_raw(path: Path, lists: list[dict]) -> dict:
    """What SD6's 'raw data' sheet holds against the PRMT5 lists (it is not the export they were drawn from)."""
    groups, counts = read_export(path, SD6_CHECK_SHEET)
    by_seq = defaultdict(list)
    for g in groups:
        by_seq[g['peptide']].append(g)
    on = lambda r, g: r['accession'] in [a.strip() for a in g['master_accessions'].split(';')]  # noqa: E731
    found = {(r['accession'], r['peptide']) for r in lists if any(on(r, g) for g in by_seq.get(r['peptide'], []))}
    same = sum(1 for r in lists if any(on(r, g) and g['ratio_heavy_light'] is not None
                                        and round(g['ratio_heavy_light'], 6) == round(r['ratio'], 6) for g in by_seq.get(r['peptide'], [])))
    pairs = {(r['accession'], r['peptide']) for r in lists}
    return {'sheet': SD6_CHECK_SHEET, 'master_proteins': counts.get('master_protein_rows', 0), 'peptide_groups': counts['peptide_groups'],
            'methyl_arginine_groups': counts['methyl_arginine_groups'],
            'listed_peptides': len(pairs), 'listed_peptides_with_a_methyl_arginine_group_here': len(found),
            'list_rows': len(lists), 'list_rows_whose_ratio_a_group_here_has': same,
            'conclusion': 'not the export the PRMT5 lists come from; not used'}


def fasta_genes(path: Path) -> dict[str, str]:
    out = {}
    with open(path, encoding='utf-8') as fh:
        for line in fh:
            if line.startswith('>'):
                m = re.search(r'\bGN=(\S+)', line)
                out[line[1:].split('|')[1]] = m.group(1) if m else ''
    return out


# ---- placing ----------------------------------------------------------------------------------------------------------

def reported_starts(value: str, accession: str) -> list[int]:
    """Starts that 'Positions in Master Proteins' gives on one master ('P18583 [1032-1039]; [1040-1047]; Q9 [5-12]')."""
    starts, current = [], None
    for token in value.split(';'):
        m = re.fullmatch(r'\s*(\S+)?\s*\[(\d+)-(\d+)\]\s*', token)
        if m:
            current = m.group(1) or current
            if current == accession:
                starts.append(int(m.group(2)))
    return starts


def place(peptide: str, accession: str, canonical: dict[str, str], given_starts: list[int]) -> tuple[str, int | None]:
    seq = canonical.get(accession)
    if seq is None:
        return 'accession_not_in_canonical_fasta', None
    hits = [m.start() + 1 for m in re.finditer(f'(?={peptide})', seq)]
    if not hits:
        return 'peptide_not_in_canonical', None
    if len(hits) > 1:
        return 'peptide_ambiguous', None
    return ('position_verified' if hits[0] in given_starts else 'placed_by_sequence'), hits[0]


def log2(x: float | None) -> float | None:
    return None if x is None or x <= 0 else math.log2(x)


def source_call(ratio: float | None, degrees: set, arm: Arm) -> str:
    if not arm.scope or ratio is None or not degrees & set(arm.scope):
        return ''
    return 'decreased' if ratio < SOURCE_CUTOFF else pc.NO_CHANGE


def uniform_call(adj_p: float | None, l2: float | None, n: int) -> str:
    if adj_p is None or l2 is None:
        return ''
    return pc.site_response([(adj_p, l2, max(n, 1))], Q_CUTOFF, LOG2_CUTOFF, strict=True)['call']


def stage_arm(arm: Arm, groups: list[dict], lists: list[dict], canonical: dict[str, str], dropped: Counter, checks: Counter) -> list[dict]:
    """The peptide rows of one comparison: its quantified methyl-arginine groups with the list rows traced to them, then the
    list rows traced to none."""
    quantified = [g for g in groups if g['ratio_heavy_light'] is not None]
    dropped['groups_not_quantified (no H/L)'] += len(groups) - len(quantified)
    by_seq = defaultdict(list)
    for g in quantified:
        by_seq[g['peptide']].append(g)
    listed = defaultdict(lambda: defaultdict(list))   # id(group) -> offset -> [list rows]
    untraced = []
    for r in lists:
        cands = [g for g in by_seq.get(r['peptide'], []) if r['accession'] in [a.strip() for a in g['master_accessions'].split(';')]
                 and round(g['ratio_heavy_light'], 6) == round(r['ratio'], 6) and g['psms'] == r['psms']]
        if len(cands) > 1:
            best = [g for g in cands if r['degree'] in marks(g['modifications'])[0].get(r['offset'], set())]
            best = best or [g for g in cands if marks(g['modifications'])[1]]
            cands = best or cands
            checks['list_rows_matching_several_groups'] += 1
            checks['list_rows_assigned_to_several_groups'] += len(cands) > 1
        if not cands:
            untraced.append(r)
            continue
        checks['list_rows_traced_to_a_group'] += 1
        for g in cands:
            listed[id(g)][r['offset']].append(r)
    checks['list_rows_not_traced_to_a_group'] += len(untraced)

    rows = []
    for g in sorted(quantified, key=lambda g: g['source_row']):
        masters = [a.strip() for a in g['master_accessions'].split(';') if a.strip()]
        loc, unloc, lys = marks(g['modifications'])
        lst = listed.get(id(g), {})
        named = {r['accession'] for v in lst.values() for r in v}
        if len(named) == 1 and masters and next(iter(named)) != masters[0]:
            masters = [next(iter(named))] + [a for a in masters if a not in named]   # a shared peptide goes where the paper put it
            checks['shared_groups_placed_on_the_master_the_lists_name'] += 1
        status, start = place(g['peptide'], masters[0], canonical, reported_starts(g['positions_in_masters'], masters[0])) if masters \
            else ('no_master_protein', None)
        offsets = sorted(set(loc) | set(lst))
        degrees = {o: (loc.get(o) or {r['degree'] for r in lst[o]}) for o in offsets}
        if any(not 1 <= o <= len(g['peptide']) or g['peptide'][o - 1] != 'R' for o in loc):
            raise SystemExit(f'{arm.comparison} row {g["source_row"]}: a localised mark is not on an R of {g["peptide"]}')
        kept = []
        if start:
            for o in offsets:
                if residue_is(canonical, masters[0], start + o - 1, 'R'):
                    kept.append(o)
                else:
                    dropped['positions_not_arginine_on_canonical'] += 1
        if status not in PLACED:
            dropped[f'groups_not_placed:{status}'] += 1
            dropped['list_rows_on_groups_not_placed'] += sum(len(v) for v in lst.values())
        elif not kept:
            dropped['placed_groups_without_a_site (marks not localised, not listed)'] += 1
        for o, rs in lst.items():
            for r in rs:
                if status in PLACED and o in kept:
                    r['_staged'] = (masters[0], start + o - 1)
                else:
                    r.setdefault('_dropped', status if status not in PLACED else 'position_not_arginine_on_canonical')
        l2 = log2(g['ratio_heavy_light'])
        all_degrees = set().union(*degrees.values()) if degrees else set()
        if unloc and not all_degrees:
            all_degrees = {DEGREE[k] for _, k, w in MOD.findall(g['modifications']) if k in DEGREE and 'R' in w}
        rows.append({
            'comparison': arm.comparison, 'source_file': SD[arm.sd][1], 'source_sheet': arm.raw_sheet, 'source_rows': str(g['source_row']),
            'peptide': g['peptide'], 'annotated_sequence': g['annotated_sequence'], 'modifications': g['modifications'],
            'master_accessions': g['master_accessions'], 'canonical_UniProtAC': masters[0] if masters and status in PLACED else '',
            'placement': status, 'canonical_start': start if status in PLACED else '',
            'reported_start': ';'.join(str(s) for s in reported_starts(g['positions_in_masters'], masters[0])) if masters else '',
            'other_master_entries': ';'.join(masters[1:]),
            'methyl_offsets': ';'.join(str(o) for o in kept),
            'methyl_positions': ';'.join(str(start + o - 1) for o in kept) if status in PLACED else '',
            'methyl_degrees': ';'.join('+'.join(sorted(degrees[o])) for o in kept),
            'localised_by': ';'.join('search' if o in loc else 'list' for o in kept),
            'n_methylarginines': len(kept) if status in PLACED else 0,
            'unlocalised_methyl_marks': unloc, 'lysine_methyl_marks': lys, 'psms': g['psms'],
            'quan_info': g['quan_info'], 'quan_usage': g['quan_usage'],
            'ratio_heavy_light': g['ratio_heavy_light'], 'log2_treated_control': l2, 'adj_p': g['adj_p'],
            'abundance_light': g['abundance_light'], 'abundance_heavy': g['abundance_heavy'],
            'abolished': g['ratio_heavy_light'] == FLOOR, 'ratio_at_ceiling': g['ratio_heavy_light'] == CEILING,
            'listed_in': ';'.join(sorted({f'{r["degree"]} {r["kind"]}' for v in lst.values() for r in v})),
            'listed_offsets': ';'.join(str(o) for o in sorted(lst)),
            'call_source': source_call(g['ratio_heavy_light'], all_degrees, arm),
            'call_uniform': uniform_call(g['adj_p'], l2, len(kept)),
            '_lists': lst, '_listed_offsets_kept': [o for o in kept if o in lst],
        })

    # list rows not traced to a quantified group: one peptide per accession, sequence, ratio and PSM count
    own = defaultdict(list)
    for r in untraced:
        own[(r['accession'], r['peptide'], round(r['ratio'], 6), r['psms'])].append(r)
    for (acc, peptide, ratio, psms), rs in sorted(own.items(), key=lambda kv: (kv[1][0]['sheet'], kv[1][0]['row'])):
        base = acc.split('-')[0]
        status, start = place(peptide, base, canonical, sorted({r['reported_position'] - r['offset'] + 1 for r in rs}) if '-' not in acc else [])
        offs = sorted({r['offset'] for r in rs})
        lst = {o: [r for r in rs if r['offset'] == o] for o in offs}
        kept = []
        if start:
            for o in offs:
                if residue_is(canonical, base, start + o - 1, 'R'):
                    kept.append(o)
                else:
                    dropped['list_positions_not_arginine_on_canonical'] += 1
        if status not in PLACED:
            dropped[f'list_peptides_not_placed:{status}'] += 1
            dropped['list_rows_on_peptides_not_placed'] += len(rs)
        for o in offs:
            for r in lst[o]:
                if status in PLACED and o in kept:
                    r['_staged'] = (base, start + o - 1)
                else:
                    r.setdefault('_dropped', status if status not in PLACED else 'position_not_arginine_on_canonical')
        degrees = {o: {r['degree'] for r in lst[o]} for o in offs}
        rows.append({
            'comparison': arm.comparison, 'source_file': SD[arm.sd][1], 'source_sheet': 'regulated lists',
            'source_rows': ';'.join(f'{r["degree"]} {r["kind"]}:{r["row"]}' for r in sorted(rs, key=lambda r: (r['sheet'], r['row']))),
            'peptide': peptide, 'annotated_sequence': '', 'modifications': '', 'master_accessions': acc,
            'canonical_UniProtAC': base if status in PLACED else '', 'placement': status, 'canonical_start': start if status in PLACED else '',
            'reported_start': ';'.join(str(s) for s in sorted({r['reported_position'] - r['offset'] + 1 for r in rs})), 'other_master_entries': '',
            'methyl_offsets': ';'.join(str(o) for o in kept), 'methyl_positions': ';'.join(str(start + o - 1) for o in kept) if status in PLACED else '',
            'methyl_degrees': ';'.join('+'.join(sorted(degrees[o])) for o in kept), 'localised_by': ';'.join('list' for _ in kept),
            'n_methylarginines': len(kept) if status in PLACED else 0, 'unlocalised_methyl_marks': '', 'lysine_methyl_marks': '', 'psms': psms,
            'quan_info': '', 'quan_usage': '', 'ratio_heavy_light': ratio, 'log2_treated_control': log2(ratio), 'adj_p': None,
            'abundance_light': None, 'abundance_heavy': None, 'abolished': ratio == FLOOR, 'ratio_at_ceiling': ratio == CEILING,
            'listed_in': ';'.join(sorted({f'{r["degree"]} {r["kind"]}' for r in rs})), 'listed_offsets': ';'.join(str(o) for o in offs),
            'call_source': source_call(ratio, set().union(*degrees.values()), arm), 'call_uniform': '',
            '_lists': lst, '_listed_offsets_kept': kept,
        })
    return rows


def site_rows(arm: Arm, peptides: list[dict], canonical: dict[str, str], genes: dict[str, str]) -> list[dict]:
    groups = defaultdict(list)
    for p in peptides:
        if p['placement'] not in PLACED or not p['methyl_positions']:
            continue
        for o, pos, deg, how in zip(p['methyl_offsets'].split(';'), p['methyl_positions'].split(';'), p['methyl_degrees'].split(';'),
                                    p['localised_by'].split(';')):
            groups[(p['canonical_UniProtAC'], int(pos))].append((p, int(o), deg, how))
    out = []
    for (acc, pos), vals in sorted(groups.items()):
        peps = [v[0] for v in vals]
        lst = [r for p, o, _, _ in vals for r in p['_lists'].get(o, [])]
        lst = sorted({(r['sheet'], r['row']): r for r in lst}.values(), key=lambda r: (r['ratio'], r['sheet'], r['row']))
        src = [p['call_source'] for p in peps if p['call_source']]
        call = 'decreased' if lst else (pc.NO_CHANGE if src else '')
        res = pc.site_response([(p['adj_p'], p['log2_treated_control'], p['n_methylarginines']) for p in peps if p['adj_p'] is not None],
                               Q_CUTOFF, LOG2_CUTOFF, strict=True)
        l2 = sorted(p['log2_treated_control'] for p in peps)
        mid = len(l2) // 2
        out.append({
            'comparison': arm.comparison, 'canonical_UniProtAC': acc, 'position': pos, 'residue': canonical[acc][pos - 1],
            'gene': genes.get(acc, ''), 'cell_line': CELL_LINE, 'enzyme': arm.enzyme, 'perturbation': arm.perturbation,
            'degrees': ';'.join(d for d in ('Rme1', 'Rme2') if any(d in v[2] for v in vals)),
            'localised_by_search': any(v[3] == 'search' for v in vals) if any(p['annotated_sequence'] for p in peps) else '',
            'n_peptides': len(peps), 'n_multi_site_peptides': sum(p['n_methylarginines'] > 1 for p in peps),
            'log2_median': (l2[mid] if len(l2) % 2 else (l2[mid - 1] + l2[mid]) / 2),
            'log2_min': l2[0], 'ratio_min': min(p['ratio_heavy_light'] for p in peps),
            'call_source': call, 'abolished_source': any(r['kind'] == 'abolished' for r in lst),
            'listed_in': ';'.join(sorted({f'{r["degree"]} {r["kind"]}' for r in lst})),
            'reported_positions': ';'.join(str(x) for x in sorted({r['reported_position'] for r in lst} - {pos})),
            'source_ratio': lst[0]['ratio'] if lst else None, 'source_protein_ratio': lst[0]['protein_ratio'] if lst else None,
            'source_normalised_ratio': lst[0]['normalised_ratio'] if lst else None,
            'call_source_rule_peptides': pc.peptide_call(src.count('decreased'), 0) if src else '',
            'n_peptides_with_adj_p': res['n_peptides'], 'adj_p_min': res['adj_p_min'],
            'n_decreased': res['n_decreased'] if res['n_peptides'] else '', 'n_increased': res['n_increased'] if res['n_peptides'] else '',
            'call_uniform': res['call'], 'call_uniform_single_site_peptides': res['call_single_site_peptides'],
            'placement': ';'.join(sorted({p['placement'] for p in peps})),
        })
    return out


# ---- checks -----------------------------------------------------------------------------------------------------------

def group_key(p: dict) -> str:
    """A peptide group across exports: sequence and methyl marks (labels, oxidation and carbamidomethyl left out)."""
    m = '; '.join(sorted(f'{n}{k}[{w}]' for n, k, w in MOD.findall(p['modifications'])))
    return f'{p["peptide"]}|{m}'


def concordance(a: list[dict], b: list[dict]) -> dict:
    """Peptide groups quantified in both exports: how many changed at least twofold the same way and the opposite way."""
    la = {group_key(p): p['log2_treated_control'] for p in a if p['annotated_sequence']}
    lb = {group_key(p): p['log2_treated_control'] for p in b if p['annotated_sequence']}
    both = sorted(set(la) & set(lb))
    x = pd.Series([la[k] for k in both])
    y = pd.Series([lb[k] for k in both])
    same = int(((x <= -1) & (y <= -1)).sum() + ((x >= 1) & (y >= 1)).sum())
    opposite = int(((x <= -1) & (y >= 1)).sum() + ((x >= 1) & (y <= -1)).sum())
    return {'groups_in_both': len(both), 'spearman': round(float(x.corr(y, method='spearman')), 3) if len(both) > 2 else None,
            'both_decreased': int(((x <= -1) & (y <= -1)).sum()), 'both_increased': int(((x >= 1) & (y >= 1)).sum()),
            'twofold_same_direction': same, 'twofold_opposite_direction': opposite}


def manifest_inputs(manifest: Path | None, paths: dict[str, Path]) -> dict[str, dict]:
    """Each input's row in inputs/MANIFEST.tsv (sources/raw/li2021/<name>), checked by MD5."""
    if manifest is None:
        return {}
    table = pd.read_csv(manifest, sep='\t', dtype=str, keep_default_na=False)
    rows = {}
    for p in paths.values():
        rel = f'sources/raw/li2021/{p.name}'
        hit = table[table['file'] == rel]
        if len(hit) != 1:
            raise SystemExit(f'{manifest}: no single row {rel}')
        row = hit.iloc[0].to_dict()
        if row['md5'] != md5_file(p):
            raise SystemExit(f'{p}: MD5 {md5_file(p)} differs from its MANIFEST row ({row["md5"]})')
        rows[p.name] = row
    return rows


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, 'rb') as fh:
        for block in iter(lambda: fh.read(1 << 20), b''):
            h.update(block)
    return h.hexdigest()


def write_rows(outdir: Path, paths: dict[str, Path], md5s: dict[str, str], inputs_rows: dict[str, dict]) -> None:
    derived = '; '.join(f'{p.name} (MD5 {md5s[p.name]})' for p in paths.values())
    source = '; '.join(ESM + p.name for p in paths.values())
    cols = ['file', 'bytes', 'md5', 'rows', 'redistribution', 'content', 'source', 'version_or_date', 'citation', 'licence',
            'derived_from', 'how_made_or_checked', 'original_name']
    rows = []
    for name, content in MANIFEST_ROWS.items():
        path = outdir / name
        n = '' if name.endswith('.json') else sum(1 for _ in open(path, encoding='utf-8')) - 1
        rows.append({'file': f'layers/{name}', 'bytes': path.stat().st_size, 'md5': md5_file(path), 'rows': n, 'redistribution': 'git',
                     'content': content, 'source': source, 'version_or_date': 'Supplementary Data 1, 3, 5 and 6 as published',
                     'citation': CITATION, 'licence': LICENCE, 'derived_from': derived,
                     'how_made_or_checked': HOW if not name.endswith('.json') else 'scripts/50_stage_li2021_perturbations.py',
                     'original_name': ''})
    pd.DataFrame(rows, columns=cols).to_csv(outdir / 'manifest_rows.tsv', sep='\t', index=False, lineterminator='\n')
    raw = []
    for key, p in paths.items():
        row = inputs_rows.get(p.name, {})
        date = re.search(r'Downloaded (\d{4}-\d\d-\d\d)', row.get('how_made_or_checked', ''))
        raw.append({'file': p.name, 'url': ESM + p.name, 'bytes': p.stat().st_size, 'md5': md5s[p.name], 'sha256': sha256_file(p),
                    'downloaded': date.group(1) if date else '', 'version': f'{SD[key][1]}, as published',
                    'kept': f'zenodo, as sources/raw/li2021/{p.name} (its inputs/MANIFEST.tsv row)'})
    pd.DataFrame(raw, columns=['file', 'url', 'bytes', 'md5', 'sha256', 'downloaded', 'version', 'kept']).to_csv(
        outdir / 'raw_files_rows.tsv', sep='\t', index=False, lineterminator='\n')


# ---- main -------------------------------------------------------------------------------------------------------------

PEPTIDE_COLUMNS = ['comparison', 'source_file', 'source_sheet', 'source_rows', 'peptide', 'annotated_sequence', 'modifications',
                   'master_accessions', 'canonical_UniProtAC', 'placement', 'canonical_start', 'reported_start', 'other_master_entries',
                   'methyl_offsets', 'methyl_positions', 'methyl_degrees', 'localised_by', 'n_methylarginines', 'unlocalised_methyl_marks',
                   'lysine_methyl_marks', 'psms', 'quan_info', 'quan_usage', 'ratio_heavy_light', 'log2_treated_control', 'adj_p',
                   'abundance_light', 'abundance_heavy', 'abolished', 'ratio_at_ceiling', 'listed_in', 'listed_offsets', 'call_source',
                   'call_uniform']


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    for key, (name, label) in SD.items():
        ap.add_argument(f'--{key}', required=True, help=f'Li 2021 {label} ({name})')
    ap.add_argument('--fasta', default='inputs/reference/uniprot_human_reviewed_2026_03.fasta')
    ap.add_argument('--manifest', default='', help='inputs/MANIFEST.tsv, read only: checks each input against its sources/raw/li2021 row')
    ap.add_argument('--outdir', required=True)
    args = ap.parse_args()
    paths = {key: Path(getattr(args, key)) for key in SD}
    for key, p in paths.items():
        if p.name != SD[key][0]:
            raise SystemExit(f'--{key} should be {SD[key][0]}, not {p.name}')
    fasta, outdir = Path(args.fasta), Path(args.outdir)
    md5s = {p.name: md5_file(p) for p in paths.values()}
    inputs_rows = manifest_inputs(Path(args.manifest) if args.manifest else None, paths)
    canonical, genes = parse_fasta(fasta), fasta_genes(fasta)

    peptides, sites, per_arm, dropped_all, lists_by_arm = [], [], {}, {}, {}
    exports = {}
    for arm in ARMS:
        lists = read_lists(paths[arm.sd], arm) if arm.lists else []
        lists_by_arm[arm.comparison] = lists
        groups, counts = read_export(paths[arm.sd], arm.raw_sheet) if arm.raw_sheet else ([], {})
        dropped, checks = Counter(), Counter()
        rows = stage_arm(arm, groups, lists, canonical, dropped, checks)
        srows = site_rows(arm, rows, canonical, genes)
        exports[arm.comparison] = rows
        peptides += rows
        sites += srows
        dropped_all[arm.comparison] = dict(sorted(dropped.items()))
        listed_sites = {(r['accession'], r['reported_position']) for r in lists}
        abolished_sites = {(r['accession'], r['reported_position']) for r in lists if r['kind'] == 'abolished'}
        s = pd.DataFrame(srows)
        dec = s[s['call_source'] == 'decreased'] if len(s) else s
        rule_dec = s[s['call_source_rule_peptides'] == 'decreased'] if len(s) else s
        # where each listed site (accession and reported position) went
        outcome, rank = {}, {'staged_at_the_reported_position': 0, 'staged_at_another_position (the canonical sequence differs from '
                                                                       'the one searched)': 1,
                             'staged_on_the_other_entry_of_a_shared_peptide (the lists name both; the peptide is placed once)': 2}
        for r in lists:
            k = (r['accession'], r['reported_position'])
            if '_staged' in r:
                o = list(rank)[0 if r['_staged'] == (r['accession'].split('-')[0], r['reported_position'])
                               else 1 if r['_staged'][0] == r['accession'].split('-')[0] else 2]
            else:
                o = f'not_staged:{r.get("_dropped", "not traced")}'
            if k not in outcome or rank.get(o, 9) < rank.get(outcome[k], 9):
                outcome[k] = o
        remapped = sorted(f'{r["accession"]} R{r["reported_position"]} -> R{r["_staged"][1]}' for r in lists if '_staged' in r
                          and r['_staged'][0] == r['accession'].split('-')[0] and r['_staged'][1] != r['reported_position'])
        other_entry = sorted(f'{r["accession"]} R{r["reported_position"]} ({r["gene"]}) -> {r["_staged"][0]} R{r["_staged"][1]}' for r in lists
                             if '_staged' in r and r['_staged'][0] != r['accession'].split('-')[0])
        not_staged = sorted({f'{r["accession"]} R{r["reported_position"]} ({r["gene"]} {r["peptide"]}): {outcome[(r["accession"], r["reported_position"])]}'
                             for r in lists if outcome[(r['accession'], r['reported_position'])].startswith('not_staged')})
        norm = [r for r in lists if r['kind'] == 'fc2' and r['normalised_ratio'] > SOURCE_CUTOFF]
        arith = sum(1 for r in lists if abs(r['normalised_ratio'] - r['ratio'] / r['protein_ratio']) > 1e-6 * max(1, r['normalised_ratio']))
        fc2 = {(r['accession'], r['reported_position']) for r in lists if r['kind'] == 'fc2'}
        per_arm[arm.comparison] = {
            'source': f'{SD[arm.sd][1]} ({paths[arm.sd].name})', 'export_sheet': arm.raw_sheet, 'lists': [s_ for s_, _, _ in arm.lists],
            'perturbation': arm.perturbation, 'enzyme': arm.enzyme, 'kind': arm.kind, 'degrees_called_by_the_paper': list(arm.scope),
            'note': arm.note, 'export': counts,
            'peptide_rows': len(rows), 'peptide_rows_placed': sum(r['placement'] in PLACED for r in rows),
            'placement': dict(sorted(Counter(r['placement'] for r in rows).items())),
            'peptide_call_source': dict(sorted(Counter(r['call_source'] or '(none)' for r in rows).items())),
            'peptide_call_uniform': dict(sorted(Counter(r['call_uniform'] or '(no adjusted p)' for r in rows).items())),
            'peptides_abolished': sum(bool(r['abolished']) for r in rows),
            'list_rows': {sh: sum(r['sheet'] == sh for r in lists) for sh, _, _ in arm.lists},
            'listed_sites_as_reported': {'fc2_or_abolished': len(listed_sites), 'proteins': len({a for a, _ in listed_sites}),
                                         'abolished': len(abolished_sites), 'abolished_proteins': len({a for a, _ in abolished_sites}),
                                         'abolished_not_in_fc2_lists': len(abolished_sites - fc2)},
            'listed_sites_outcome': dict(sorted(Counter(outcome.values()).items())),
            'listed_sites_remapped': sorted(set(remapped)), 'listed_sites_on_the_other_entry': sorted(set(other_entry)),
            'listed_sites_not_staged': not_staged,
            'list_fc2_rows_with_normalised_ratio_above_0.5': {'rows': len(norm), 'proteins': len({r['accession'] for r in norm})},
            'list_rows_where_normalised_differs_from_ratio_over_protein_ratio': arith,
            'checks': dict(sorted(checks.items())),
            'sites': len(srows), 'site_proteins': int(s['canonical_UniProtAC'].nunique()) if len(s) else 0,
            'sites_by_degree': dict(sorted(Counter(s['degrees']).items())) if len(s) else {},
            'site_call_source': dict(sorted(Counter(s['call_source'].replace('', '(none)')).items())) if len(s) else {},
            'sites_decreased_source': len(dec), 'proteins_decreased_source': int(dec['canonical_UniProtAC'].nunique()) if len(dec) else 0,
            'sites_abolished_source': int(s['abolished_source'].sum()) if len(s) else 0,
            'proteins_abolished_source': int(s.loc[s['abolished_source'], 'canonical_UniProtAC'].nunique()) if len(s) else 0,
            'sites_quantified_in_scope': int((s['call_source'] != '').sum()) if len(s) else 0,
            'sites_decreased_not_localised_by_search': int((dec['localised_by_search'] == False).sum()) if len(dec) else 0,  # noqa: E712
            'site_call_source_rule_peptides': dict(sorted(Counter(s['call_source_rule_peptides'].replace('', '(none)')).items())) if len(s) else {},
            'sites_rule_decreased_not_listed': int(((s['call_source_rule_peptides'] == 'decreased') & (s['call_source'] != 'decreased')).sum()) if len(s) else 0,
            'sites_listed_rule_not_decreased': int(((s['call_source'] == 'decreased') & (s['call_source_rule_peptides'] != 'decreased')).sum()) if len(s) else 0,
            'site_call_uniform': dict(sorted(Counter(s['call_uniform'].replace('', '(no adjusted p)')).items())) if len(s) else {},
            'site_call_uniform_single_site_peptides': dict(sorted(Counter(s['call_uniform_single_site_peptides'].replace('', '(none)')).items())) if len(s) else {},
            'paper': PAPER.get(arm.comparison, {}),
        }
        if len(rule_dec):
            per_arm[arm.comparison]['proteins_rule_decreased'] = int(rule_dec['canonical_UniProtAC'].nunique())

    # the paper's way of counting quantified sites: every arginine of a peptide whose mark PD did not localise
    for arm in ARMS:
        if not arm.raw_sheet:
            continue
        cand = set()
        for r in exports[arm.comparison]:
            if r['placement'] not in PLACED or not r['annotated_sequence']:
                continue
            loc, unloc, _ = marks(r['modifications'])
            deg = {d for v in loc.values() for d in v} | ({DEGREE[k] for _, k, w in MOD.findall(r['modifications']) if k in DEGREE and 'R' in w} if unloc else set())
            if arm.scope and not deg & set(arm.scope):
                continue
            offs = set(loc) | ({i + 1 for i, ch in enumerate(r['peptide']) if ch == 'R'} if unloc else set())
            cand |= {(r['canonical_UniProtAC'], int(r['canonical_start']) + o - 1) for o in offs}
        per_arm[arm.comparison]['sites_quantified_counting_every_arginine_of_unlocalised_peptides'] = len(cand)

    pep = pd.DataFrame(peptides)
    if (pep['comparison'] + '|' + pep['source_rows'] + '|' + pep['peptide'] + '|' + pep['modifications']).duplicated().any():
        raise SystemExit('duplicate peptide rows')
    st = pd.DataFrame(sites)
    if st['residue'].ne('R').any():
        raise SystemExit('placed sites that are not arginines on the canonical sequence')
    save_table(pep[PEPTIDE_COLUMNS], outdir / 'li2021_peptides.tsv')
    save_table(st, outdir / 'li2021_sites.tsv')

    sd6_check = check_sd6_raw(paths['sd6'], lists_by_arm['PRMT5_si'])
    adj_floor = {c: sorted({r['adj_p'] for r in exports[c] if r['abolished'] and r['adj_p'] is not None}) for c in ('PRMT7_si', 'PRMT7_si_rep2', 'PRMT7_SGC3027')}
    boundary = [f'{c} {r["master_accessions"]} {r["peptide"]} {r["modifications"]} H/L={r["ratio_heavy_light"]:g} listed={bool(r["listed_in"])}'
                for c in exports for r in exports[c] if r['ratio_heavy_light'] == SOURCE_CUTOFF and r['annotated_sequence']]
    summary = {
        'study': STUDY, 'doi': DOI, 'licence': LICENCE, 'pride': 'PXD022424 (raw MS files, not used)', 'cell_line': CELL_LINE,
        'inputs': {p.name: {'md5': md5s[p.name], 'label': SD[k][1], 'manifest_row': f'sources/raw/li2021/{p.name}' if p.name in inputs_rows else '',
                            'sheets_read': sorted({a.raw_sheet for a in ARMS if a.sd == k and a.raw_sheet} | {s_ for a in ARMS if a.sd == k for s_, _, _ in a.lists}
                                                  | ({SD6_CHECK_SHEET} if k == 'sd6' else set()))}
                   for k, p in paths.items()},
        'canonical_fasta': {'file': fasta.name, 'md5': md5_file(fasta), 'sequences': len(canonical)},
        'orientation': ORIENTATION, 'log2_treated_control': 'log2 of the export\'s Abundance Ratio (Heavy)/(Light) (of the listed ratio for list-only peptides)',
        'normalisation': NORMALISATION, 'replicates': REPLICATES,
        'adjusted_p': ('Proteome Discoverer 2.2 "Abundance Ratio Adj. P-Value: (Heavy) / (Light)" of SD1 and SD3, as given; the paper does not '
                       'describe it. Each export compares one SILAC pair of cultures (SD1 repeat 1 enters its 15 raw files as 15 sample '
                       'pairs), so it is not a test over biological replicates. SD5 gives none. Every peptide group at the 0.01 floor (and '
                       'at the 100 ceiling) has the same adjusted p within an export, so every abolished peptide passes the uniform rule: '
                       + '; '.join(f'{c} {v}' for c, v in adj_floor.items())),
        'paper_rule': PAPER_RULE, 'source_rule': SOURCE_RULE, 'uniform_rule': UNIFORM_RULE,
        'boundary': {'groups_at_H/L_0.5': boundary,
                     'reading': 'sheet names say FC>2 and the one group at exactly 0.5 whose site is listed by no other group is not listed: H/L < 0.5'},
        'comparisons': per_arm,
        'dropped': dropped_all,
        'sgc3027_orientation_check': {
            'question': 'is the SGC3027 H/L treated/untreated? (the Methods sentence read literally says untreated/treated)',
            'with_PRMT7_si': concordance(exports['PRMT7_si'], exports['PRMT7_SGC3027']),
            'with_PRMT7_si_rep2': concordance(exports['PRMT7_si_rep2'], exports['PRMT7_SGC3027']),
            'reading': 'more groups change twofold the same way as with siPRMT7 than the opposite way: heavy = SGC3027, as the paper reads it'},
        'prmt7_repeats': {'repeat1_vs_repeat2': concordance(exports['PRMT7_si'], exports['PRMT7_si_rep2']),
                          'note': 'the two siPRMT7 exports agree weakly; the paper\'s lists use repeat 1 only'},
        'sd6_raw_data_check': sd6_check,
        'shared_peptides_placed_on_first_master': int(((pep['placement'].isin(PLACED)) & (pep['other_master_entries'] != '')).sum()),
    }
    (outdir / 'li2021_summary.json').write_text(json.dumps(summary, indent=2) + '\n', encoding='utf-8')
    write_rows(outdir, paths, md5s, inputs_rows)
    print(json.dumps({c: {k: v for k, v in d.items() if k in ('peptide_rows', 'sites', 'sites_decreased_source', 'proteins_decreased_source',
                                                               'sites_abolished_source', 'sites_quantified_in_scope', 'site_call_uniform')}
                      for c, d in per_arm.items()}, indent=2))


if __name__ == '__main__':
    main()
