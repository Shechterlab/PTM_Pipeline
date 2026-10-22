#!/usr/bin/env python3
"""Stage the measured-peptide backgrounds of CAPRI (Panhale et al. 2019) and RBDmap (Castello et al. 2016) from PRIDE (inputs/interactions/).

    python -I scripts/43f_stage_peptide_backgrounds.py [--only capri|rbdmap] [--cache PTM_data/external] [--out inputs/interactions] [--preview [--preview-out FILE]]
    python -I scripts/43f_stage_peptide_backgrounds.py --no-stage --preview [--sites releases/v1.2.0/human_methylarginine_sites_v1.2.0.tsv]

Both experiments call a region RNA-bound from peptides measured by mass spectrometry, so a region can be RNA-bound only
where one of its peptides was measured, and methylarginines are themselves found in easily measured peptides. These
tables list every peptide each experiment measured, on the canonical sequence, with the authors' RNA-bound call: the
background a test needs to compare methylarginines with the other arginines the experiment could see (as scripts/37 does
for pCLAP with rna_pclap_cov).

CAPRI (HEK293, 254 nm UV; Panhale et al., Nat Commun 10:2682, 2019; PRIDE PXD013338, CC0). RNA-protein complexes are
captured on oligo(dT) beads and digested with Lys-C; the Lys-C peptides crosslinked to RNA stay on the RNA, are digested
with trypsin, and the tryptic peptides next to the crosslink are measured label-free. The MaxQuant search of the human
adjacent-peptide runs (20160811_human_CAPRI_Adjacent_MaxQuant_augs2016.zip: 20 runs, pools A and B of 3 UV-crosslinked
and 3 non-crosslinked biological replicates, grouped by MaxQuant into six experiments) gives peptides.txt.
  * +UV and -UV: the experiment of each run in summary.txt, checked against the sheet "Hs CAPRI Adjacent peptides MQ" of
    PRIDE's Filenames_Experimental_Design.xlsx (UV_Biorep1-3 and noUV_Biorep1-3).
  * Detection: a replicate detects a peptide when its intensity is above 0; "by MS/MS" counts the replicates that
    identified it by MS/MS rather than by matching between runs.
  * Adjacent (RNA-bound) peptides: the authors' own, the adjacent_peptides column of capri_panhale2019_adj_peptides.tsv
    (Supplementary Data 3), matched by sequence. Their rule (Methods, "Analysis of ADJ-peptides"): contaminants removed,
    missing intensities imputed with the minimum, moderated t-test (limma) of crosslinked against non-crosslinked with a
    Benjamini-Hochberg FDR < 0.05, then a mean intensity more than eightfold higher.
  * Lys-C region: CAPRI's rule, the tryptic peptide "extended in silico to the next nearest (theoretical) Lys-C cleavage
    site (lysine) in both directions" (no proline rule, no length limit). With the adjacent peptides on the authors' entry
    it gives the span of 1,712 of the 1,713 ADJ-peptides that place on the reference (the exception, in DDX17, differs in
    sequence); the staging repeats the check on the ADJ-peptides whose adjacent peptides it places on that entry.

RBDmap (HeLa, 254 nm UV; Castello et al., Mol Cell 63:696, 2016; PRIDE PXD000883, EMBL-EBI terms of use). RBPs captured
on oligo(dT) are digested with Lys-C or Arg-C and captured again: the supernatant holds the released fragments (F2), the
eluate the RNA-bound ones (F3); both, and an input (F1, crosslinked and not), are digested with trypsin, labelled by
reductive dimethylation and measured. RBDmap_{LysC,ArgC}_peptideID.txt are the MaxQuant peptide tables, *_QUAN.txt the
intensity of each peptide in each fraction and replicate (F3cCL, F2cCL, F3noCL, F1cCL, F1noCL).
  * RBDpep rule (Supplemental Experimental Procedures, "Statistical Analysis"): the log2 intensity ratio RNA-bound (F3)
    over released (F2), normalised to the location of the left (released) mode, tested against zero over the three
    replicates with limma's moderated t-test, Benjamini-Hochberg; 1% FDR = RBDpep, 10% FDR = CandidateRBDpep. The flags
    here are the authors' calls as published in Table S1 (sheets RBDpep and CandidateRBDpep, the same file scripts/37
    reads), matched by protease and peptide sequence; the Table S1 RBDpeps are all in *_QUAN.txt.
  * Measured: every identified peptide is kept; quantified_bound_or_released counts the replicates with an F2 or F3
    intensity (a peptide seen only in the input fraction could never be called RNA-bound).
  * Fragment: "extended to the two adjacent LysC or ArgC cleavage sites". The in-silico digest that reproduces Table S1's
    fragment column (8,471 of the 8,473 rows of its three sheets that place on the reference): cut after every K (Lys-C)
    or R (Arg-C), but a piece shorter than 6 residues is not cut off and joins the next one (the protein's first piece
    counts 2 residues more, so it is cut off from 4 residues); a last piece of one residue joins the fragment before it.
    The fragment of a peptide is the stretch of digest pieces it overlaps. The crosslinked residue lies in the fragment,
    outside the identified peptide. (The two Table S1 rows it misses start at an internal methionine: another isoform.)

Each peptide is placed on the canonical sequence of inputs/reference/uniprot_human_reviewed_2026_03.fasta by exact, unique
match as scripts/43d_stage_pclap.py does it (leading razor protein, then the other accessions given, a secondary accession
followed to its entry, then the gene names); everything else is dropped and counted. Reverse hits and contaminants are
removed. Neither search allowed methylarginine as a modification.

CAPRI's table is CC0 (redistribution = git). RBDmap's comes from data PRIDE gives only under the EMBL-EBI terms of use,
so it is not redistributed (checksum_only, inputs/interactions/.gitignore); this script rebuilds it (same MD5).

--preview compares the methylarginines of a release with the other arginines of the same proteins (Mantel-Haenszel odds
ratios, one stratum per protein with at least one RNA-bound region), on all arginines and restricted to those the
experiment measured. It writes nothing in inputs/.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import importlib.util
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from functools import lru_cache
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

_spec = importlib.util.spec_from_file_location('stage_pclap', Path(__file__).with_name('43d_stage_pclap.py'))
pclap = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pclap)
log, fetch, file_hash, md5_bytes, Reference = pclap.log, pclap.fetch, pclap.file_hash, pclap.md5_bytes, pclap.Reference
mantel_haenszel, write_manifest, parameters, runs = pclap.mantel_haenszel, pclap.write_manifest, pclap.parameters, pclap.runs

PRIDE_API = 'https://www.ebi.ac.uk/pride/ws/archive/v3/projects/{acc}'
PRIDE_FTP = 'https://ftp.pride.ebi.ac.uk/pride/data/archive/{path}/{name}'

CAPRI_ACC, CAPRI_PATH = 'PXD013338', '2019/07/PXD013338'
CAPRI_ZIP = '20160811_human_CAPRI_Adjacent_MaxQuant_augs2016.zip'
CAPRI_DIR = '20160811_adj_radj_combined_augs2016_Max_txt/'
CAPRI_MEMBERS = [CAPRI_DIR + m for m in ('peptides.txt', 'parameters.txt', 'summary.txt')]
CAPRI_DESIGN, CAPRI_DESIGN_SHEET = 'Filenames_Experimental_Design.xlsx', 'Hs CAPRI Adjacent peptides MQ'
CAPRI_ADJ = 'capri_panhale2019_adj_peptides.tsv'
CAPRI_TABLE, CAPRI_RAW = 'capri_panhale2019_measured_peptides.tsv', 'raw_capri_pride_files.tsv'
CAPRI_CITATION = 'Panhale et al., Nat Commun 10:2682 (2019), doi:10.1038/s41467-019-10585-3'

RBDMAP_ACC, RBDMAP_PATH = 'PXD000883', '2016/08/PXD000883'
RBDMAP_FILES = [f'RBDmap_{e}_{k}.txt' for e in ('LysC', 'ArgC') for k in ('peptideID', 'QUAN')]
RBDMAP_S1_URL = 'https://ora.ox.ac.uk/objects/uuid:092f4366-cac9-4ed7-8283-89dc6032834a/files/ma75a36ddd2e91c194d65041242e5a683'
RBDMAP_S1, RBDMAP_S1_CACHED = 'rbdmap_castello2016_table_s1.xlsx', 'crosslinks/rbdmap_castello2016_table_s1.xlsx'
RBDMAP_TABLE, RBDMAP_RAW = 'rbdmap_castello2016_measured_peptides.tsv', 'raw_rbdmap_pride_files.tsv'
RBDMAP_CITATION = 'Castello et al., Mol Cell 63:696 (2016), doi:10.1016/j.molcel.2016.06.029'
RBDMAP_SITE = {'LysC': 'K', 'ArgC': 'R'}
# The in-silico digest behind Table S1's fragments (fitted to the table: see the docstring)
RBD_MIN_PIECE, RBD_FIRST_PIECE_BONUS, RBD_MIN_LAST_PIECE = 6, 2, 2


def tsv(path: Path, columns: list, rows: list):
    with open(path, 'w', encoding='utf-8', newline='') as fh:
        w = csv.writer(fh, delimiter='\t', lineterminator='\n')
        w.writerow(columns)
        w.writerows(rows)


def saved(p: Path) -> str:
    return dt.date.fromtimestamp(p.stat().st_mtime).isoformat()


def pride_listing(cache: Path, acc: str):
    """PRIDE's project record (licence) and its full file listing (SHA-1 of each file), kept in the cache."""
    api = PRIDE_API.format(acc=acc)
    project_json = fetch(api, cache / 'pride_project.json')
    files_json = fetch(f'{api}/files/all', cache / 'pride_files.json')
    project = json.loads(project_json.read_text(encoding='utf-8'))
    listed = {f['fileName']: f for f in json.loads(files_json.read_text(encoding='utf-8'))}
    return project, listed, [project_json, files_json]


def fetch_checked(path: str, name: str, cache: Path, listed: dict) -> Path:
    """A PRIDE file, downloaded once, its SHA-1 checked against PRIDE's listing."""
    dest = fetch(PRIDE_FTP.format(path=path, name=name), cache / name)
    sha1, want = file_hash(dest, 'sha1'), listed.get(name, {}).get('checksum')
    if sha1 != want:
        raise SystemExit(f'{name}: SHA-1 {sha1} differs from PRIDE\'s {want}')
    return dest


def lysc_region_capri(seq: str, start: int, end: int):
    """CAPRI's Lys-C region of seq[start:end] (0-based, end exclusive): extended to the nearest lysine on each side, the
    peptide's own C-terminal lysine included. Returns 0-based (start, end)."""
    a = seq.rfind('K', 0, start) + 1
    b = seq.find('K', end - 1)
    return a, len(seq) if b < 0 else b + 1


@lru_cache(maxsize=None)
def rbdmap_digest(seq: str, site: str) -> tuple:
    """RBDmap's in-silico Lys-C (site K) or Arg-C (site R) digest as Table S1's fragments show it: 0-based [start, end)."""
    frags, start = [], 0
    for i, aa in enumerate(seq):
        if aa == site and i + 1 < len(seq) and i + 1 - start + (RBD_FIRST_PIECE_BONUS if start == 0 else 0) >= RBD_MIN_PIECE:
            frags.append((start, i + 1))
            start = i + 1
    if frags and len(seq) - start < RBD_MIN_LAST_PIECE:
        frags[-1] = (frags[-1][0], len(seq))
    else:
        frags.append((start, len(seq)))
    return tuple(frags)


def rbdmap_fragment(seq: str, site: str, start: int, end: int):
    fr = [f for f in rbdmap_digest(seq, site) if f[0] < end and f[1] > start]
    return fr[0][0], fr[-1][1]


def placed_counts(rows, dropped, flag_col):
    by_reason = Counter()
    for (h, _), v in dropped.items():
        by_reason[h] += v
    return {'kept': len(rows), 'dropped': sum(dropped.values()), 'dropped_by_reason': dict(sorted(by_reason.items())),
            f'dropped_{flag_col}': sum(v for (h, b), v in dropped.items() if b)}


# ---- CAPRI ---------------------------------------------------------------------------------------------------------------

def capri_design(xlsx: Path) -> dict:
    """{raw file name without .raw: experiment label} from the human adjacent-peptide sheet of the design file."""
    rows = list(openpyxl.load_workbook(xlsx, read_only=True)[CAPRI_DESIGN_SHEET].iter_rows(values_only=True))
    if not rows or str(rows[0][0]).strip() != 'Filename':
        raise SystemExit(f'{xlsx.name}: unexpected sheet {CAPRI_DESIGN_SHEET}')
    return {re.sub(r'\.raw$', '', str(r[0]).strip()): str(r[1]).strip() for r in rows[1:] if r and r[0]}


def stage_capri(args, ref: Reference) -> tuple[dict, list]:
    cache, out = Path(args.cache) / 'capri', Path(args.out)
    project, listed, meta = pride_listing(cache, CAPRI_ACC)
    zpath = fetch_checked(CAPRI_PATH, CAPRI_ZIP, cache, listed)
    design_path = fetch_checked(CAPRI_PATH, CAPRI_DESIGN, cache, listed)
    log(f'{CAPRI_ZIP}: {zpath.stat().st_size:,} bytes; {CAPRI_DESIGN}: SHA-1 match PRIDE; PRIDE licence: {project.get("license")!r}')
    with zipfile.ZipFile(zpath) as z:
        members = {}
        for name in CAPRI_MEMBERS:
            info = z.getinfo(name)
            if info.is_dir():
                raise SystemExit(f'{name}: not a file in {CAPRI_ZIP}')
            members[name] = z.read(info)
    par = parameters(members[CAPRI_DIR + 'parameters.txt'])
    rr = runs(members[CAPRI_DIR + 'summary.txt'])
    var_mods = sorted({m for r in rr for m in r['Variable modifications'].split(';') if m})
    methyl_r = [m for m in var_mods if re.search(r'methyl', m, re.I)]

    # +UV and -UV: MaxQuant's experiment of each run (uv_biorepN, nonuv_biorepN), checked against the authors' design
    # sheet (UV_BiorepN, noUV_BiorepN)
    design = capri_design(design_path)
    exp_of = {r['Raw file']: r['Experiment'] for r in rr}

    def uv_rep(label: str, pattern: str):
        m = re.fullmatch(pattern, label, re.I)
        return (m.group(1) is None, int(m.group(2))) if m else None

    mq = {f: uv_rep(e, r'(non)?uv_biorep(\d+)') for f, e in exp_of.items()}
    des = {f: uv_rep(lab, r'(no)?UV_Biorep(\d+)') for f, lab in design.items()}
    if set(des) != set(mq) or None in mq.values() or any(des[f] != mq[f] for f in mq):
        raise SystemExit(f'runs of summary.txt and {CAPRI_DESIGN} disagree')
    uv = sorted({exp_of[f] for f in exp_of if des[f][0]})
    nonuv = sorted({exp_of[f] for f in exp_of if not des[f][0]})
    if len(uv) != 3 or len(nonuv) != 3 or len(exp_of) != 20:
        raise SystemExit(f'unexpected experiments: {sorted(set(exp_of.values()))}')
    log(f'MaxQuant {par.get("Version")}; {len(exp_of)} runs; +UV {uv}; -UV {nonuv}; enzyme {sorted({r["Enzyme"] for r in rr})}; '
        f'variable modifications {var_mods}; methylarginine searched: {bool(methyl_r)}')

    pep = pd.read_csv(io.BytesIO(members[CAPRI_DIR + 'peptides.txt']), sep='\t', dtype=str, keep_default_na=False)
    n_all = len(pep)
    pep = pep[(pep['Reverse'] != '+') & (pep['Potential contaminant'] != '+')].reset_index(drop=True)
    inten = pep[[f'Intensity {e}' for e in uv + nonuv]].apply(pd.to_numeric).to_numpy(float)
    idt = pep[[f'Identification type {e}' for e in uv + nonuv]].to_numpy(str)

    with open(Path(args.capri_adj), encoding='utf-8') as fh:
        adj_rows = list(csv.DictReader(fh, delimiter='\t'))
    adjacent = {p for r in adj_rows for p in r['adjacent_peptides'].split('; ') if p}
    in_search = adjacent & set(pep['Sequence'])

    rows, dropped, how_count = [], Counter(), Counter()
    for i, r in pep.iterrows():
        seq, is_adj = r['Sequence'], r['Sequence'] in adjacent
        acc, s, how = ref.place(seq, r['Leading razor protein'], [p for p in r['Proteins'].split(';') if p],
                                [g for g in r['Gene names'].split(';') if g])
        if acc is None:
            dropped[(how, is_adj)] += 1
            continue
        how_count[how] += 1
        la, lb = lysc_region_capri(ref.seq[acc], s, s + len(seq))
        rows.append([acc, ref.gene.get(acc, ''), s + 1, s + len(seq), seq, la + 1, lb, how, 'true' if is_adj else 'false',
                     int((inten[i, :3] > 0).sum()), int((inten[i, 3:] > 0).sum()),
                     int((idt[i, :3] == 'By MS/MS').sum()), int((idt[i, 3:] == 'By MS/MS').sum())])
    rows.sort(key=lambda x: (x[0], x[2], x[3], x[4]))
    columns = ['uniprot_id', 'gene', 'start', 'end', 'sequence', 'lysc_start', 'lysc_end', 'mapped_by', 'adjacent_peptide',
               'uv_replicates_detected', 'nonuv_replicates_detected', 'uv_replicates_msms', 'nonuv_replicates_msms']
    out.mkdir(parents=True, exist_ok=True)
    path = out / CAPRI_TABLE
    tsv(path, columns, rows)

    # Check: each ADJ-peptide of the authors (placed on their entry as scripts/37 places it) against the Lys-C regions of
    # its adjacent peptides here. The razor protein can be another entry than the authors' (a paralog sharing the peptide).
    regions, entry = defaultdict(set), {}
    for r in rows:
        entry[r[4]] = r[0]
        if r[8] == 'true':
            regions[r[0]].add((r[5], r[6]))
    adj_check = Counter()
    for r in adj_rows:
        a = r['uniprot_id'].split('-')[0]
        accs = {a} if a in ref.seq else (ref.secondary.get(a) or ref.by_gene.get(r['gene'].upper()) or set())
        hits = [(x, ref.unique_start(x, r['adj_peptide'])) for x in accs]
        hits = [(x, i) for x, i in hits if i is not None]
        if len(hits) != 1:
            adj_check['not_placed_on_reference'] += 1
            continue
        x, i = hits[0]
        if x not in {entry.get(p) for p in r['adjacent_peptides'].split('; ')}:
            adj_check['adjacent_peptides_on_another_entry'] += 1
        else:
            adj_check['span_reproduced' if (i + 1, i + len(r['adj_peptide'])) in regions[x] else 'span_differs'] += 1
    adj_uv = Counter((r[9], r[10]) for r in rows if r[8] == 'true')
    counts = placed_counts(rows, dropped, 'adjacent')
    summary = {
        'peptides_in_file': n_all, 'measured': len(pep), **counts, 'placed_by': dict(how_count),
        'kept_adjacent': sum(r[8] == 'true' for r in rows), 'proteins': len({r[0] for r in rows}),
        'proteins_adjacent': len({r[0] for r in rows if r[8] == 'true'}),
        'adjacent_peptides_listed': len(adjacent), 'adjacent_peptides_in_search': len(in_search),
        'adj_peptides_check': dict(adj_check),
        'adjacent_by_uv_nonuv_replicates_detected': {f'{u}/{n}': v for (u, n), v in sorted(adj_uv.items())},
        'never_quantified': int(sum(r[9] + r[10] == 0 for r in rows)),
        'variable_modifications': var_mods, 'methylarginine_searched': bool(methyl_r), 'pride_licence': project.get('license'),
        'maxquant_version': par.get('Version')}
    log(f'{CAPRI_TABLE}: {len(rows):,} peptides on {summary["proteins"]:,} proteins, {summary["kept_adjacent"]:,} adjacent on '
        f'{summary["proteins_adjacent"]:,}; {path.stat().st_size / 1e6:.2f} MB')

    raw = [[CAPRI_ZIP, zpath.stat().st_size, file_hash(zpath), saved(zpath)]]
    raw += [[f'{CAPRI_ZIP}:{m}', len(b), md5_bytes(b), saved(zpath)] for m, b in members.items()]
    raw += [[p.name, p.stat().st_size, file_hash(p), saved(p)] for p in [design_path] + meta]
    raw_path = out / CAPRI_RAW
    tsv(raw_path, ['path', 'bytes', 'md5', 'saved'], raw)
    day = saved(zpath)
    search_day = dt.datetime.strptime(par.get('Date of writing', '').split()[0], '%m/%d/%Y').date().isoformat()
    n_adj_dropped = counts['dropped_adjacent']
    manifest = [
        {'file': f'{out.name}/{CAPRI_TABLE}', 'bytes': path.stat().st_size, 'md5': file_hash(path), 'rows': len(rows),
         'redistribution': 'git',
         'content': f'CAPRI (HEK293, 254 nm UV; Lys-C on the oligo(dT)-captured RBPs, trypsin on the Lys-C peptides left on the RNA) '
                    f'human peptides measured in the adjacent-peptide runs (3 UV-crosslinked and 3 non-crosslinked biological '
                    f'replicates, pools A and B), from the PRIDE MaxQuant search: each peptide on the canonical sequence (exact, '
                    f'unique match) with its Lys-C region (extended to the nearest lysine on each side, CAPRI\'s rule), whether it '
                    f'is one of the authors\' adjacent peptides (Supplementary Data 3, by sequence), and how many +UV and -UV '
                    f'replicates detected it (intensity) and identified it by MS/MS. {summary["kept_adjacent"]:,} adjacent peptides on '
                    f'{summary["proteins_adjacent"]:,} proteins; the others are the measured background. The search did not allow '
                    f'methylarginine',
         'source': f'{PRIDE_FTP.format(path=CAPRI_PATH, name=CAPRI_ZIP)} ({CAPRI_DIR}peptides.txt)',
         'version_or_date': f'PRIDE {CAPRI_ACC} (public {project.get("publicationDate")}; MaxQuant {par.get("Version")}, search '
                            f'written {search_day}), downloaded {day}',
         'citation': CAPRI_CITATION,
         'licence': f'{project.get("license")} (PRIDE {CAPRI_ACC})',
         'derived_from': f'{CAPRI_ZIP}: peptides.txt, parameters.txt, summary.txt; {CAPRI_DESIGN} ({CAPRI_RAW}); '
                         f'interactions/{CAPRI_ADJ} (adjacent peptides)',
         'how_made_or_checked': f'scripts/43f_stage_peptide_backgrounds.py (archive and design file SHA-1 checked against PRIDE; '
                                f'+UV/-UV from summary.txt checked against the design file; {len(rows):,} peptides kept, '
                                f'{counts["dropped"]:,} dropped: '
                                + ', '.join(f'{v:,} {k.replace("_", " ")}' for k, v in counts['dropped_by_reason'].items())
                                + f'; {len(in_search):,} of the {len(adjacent):,} adjacent peptides listed are in the search, '
                                  f'{n_adj_dropped:,} of them dropped). The Lys-C region rule reproduces '
                                  f'{adj_check["span_reproduced"]:,} of the {adj_check["span_reproduced"] + adj_check["span_differs"]:,} '
                                  f'ADJ-peptides whose adjacent peptides are placed on the authors\' entry; for '
                                  f'{adj_check["adjacent_peptides_on_another_entry"]:,} the leading razor protein is another entry '
                                  f'sharing the peptide (a paralog)',
         'original_name': 'peptides.txt'},
        {'file': f'{out.name}/{CAPRI_RAW}', 'bytes': raw_path.stat().st_size, 'md5': file_hash(raw_path), 'rows': len(raw),
         'redistribution': 'git',
         'content': f'The files behind {CAPRI_TABLE}: the PRIDE {CAPRI_ACC} MaxQuant output archive of the human adjacent-peptide '
                    f'runs, the three tables read from it, the experimental design file and PRIDE\'s project and file listings '
                    f'(licence; SHA-1 of each file), with their size, MD5 and date',
         'source': f'{PRIDE_FTP.format(path=CAPRI_PATH, name=CAPRI_ZIP)}; {PRIDE_FTP.format(path=CAPRI_PATH, name=CAPRI_DESIGN)}; '
                   f'{PRIDE_API.format(acc=CAPRI_ACC)}',
         'version_or_date': day, 'citation': '', 'licence': '', 'derived_from': '',
         'how_made_or_checked': 'scripts/43f_stage_peptide_backgrounds.py', 'original_name': ''}]
    return summary, manifest


# ---- RBDmap --------------------------------------------------------------------------------------------------------------

def table_s1_calls(xlsx: Path) -> dict:
    """{(protease, tryptic peptide): (sheet, fragment sequence, ProtID)} for the RBDpep (1% FDR) and CandidateRBDpep (10%
    FDR) sheets."""
    calls = {}
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    for sheet in ('RBDpep', 'CandidateRBDpep'):
        rows = wb[sheet].iter_rows(values_only=True)
        head = [str(h).strip() for h in next(rows)]
        pep_col = next(h for h in head if h.startswith('MS-identified'))
        frag_col = next(h for h in head if h.startswith('LysC/ArgC'))
        for r in rows:
            d = dict(zip(head, r))
            if d.get(pep_col):
                calls.setdefault((str(d['enzyme']), str(d[pep_col])), (sheet, str(d[frag_col] or ''), str(d['ProtID'] or '').split('-')[0]))
    return calls


def stage_rbdmap(args, ref: Reference) -> tuple[dict, list]:
    cache, out = Path(args.cache) / 'rbdmap', Path(args.out)
    project, listed, meta = pride_listing(cache, RBDMAP_ACC)
    paths = {name: fetch_checked(RBDMAP_PATH, name, cache, listed) for name in RBDMAP_FILES}
    s1 = fetch(RBDMAP_S1_URL, cache / RBDMAP_S1)
    s1_md5, s1_known = file_hash(s1), ''
    raw_cache = Path(args.raw_cache)
    if raw_cache.exists():
        with open(raw_cache, encoding='utf-8') as fh:
            s1_known = next((r['md5'] for r in csv.DictReader(fh, delimiter='\t') if r['path'] == RBDMAP_S1_CACHED), '')
    if s1_known and s1_md5 != s1_known:
        raise SystemExit(f'{RBDMAP_S1}: MD5 {s1_md5} differs from the Table S1 scripts/37 read ({s1_known}, raw_cache_files.tsv)')
    log(f'RBDmap: {len(paths)} PRIDE files, SHA-1 match PRIDE; PRIDE licence: {project.get("license")!r}; Table S1 MD5 '
        f'{"matches the one scripts/37 read" if s1_known else s1_md5}')
    calls = table_s1_calls(s1)

    rows, dropped, how_count, per_enzyme, s1_check = [], Counter(), Counter(), {}, Counter()
    for enz in ('LysC', 'ArgC'):
        site = RBDMAP_SITE[enz]
        pid = pd.read_csv(paths[f'RBDmap_{enz}_peptideID.txt'], sep='\t', dtype=str, keep_default_na=False)
        n_all = len(pid)
        pid = pid[(pid['Reverse'] != '+') & (pid['Contaminant'] != '+')].reset_index(drop=True)
        quan = pd.read_csv(paths[f'RBDmap_{enz}_QUAN.txt'], sep='\t', dtype={'Sequence': str})
        quan.columns = [c.replace('ccL', 'cCL') for c in quan.columns]
        quan = quan.set_index('Sequence')
        if quan.index.duplicated().any() or set(quan.index) != set(pid['Sequence']):
            raise SystemExit(f'RBDmap {enz}: the QUAN and peptideID sequences differ')
        reps = sorted({c.split('_')[1] for c in quan.columns if '_rep' in c})
        q = quan.loc[pid['Sequence']]
        bound = np.stack([q[f'F3cCL_{r}'].to_numpy(float) for r in reps], 1)
        released = np.stack([q[f'F2cCL_{r}'].to_numpy(float) for r in reps], 1)
        inp = np.stack([(q[f'F1cCL_{r}'].to_numpy(float) > 0) | (q[f'F1noCL_{r}'].to_numpy(float) > 0) for r in reps], 1)
        both = (bound > 0) & (released > 0)
        lr = np.where(both, np.log2(np.where(both, bound, 1) / np.where(both, released, 1)), np.nan)
        n_both = both.sum(1)
        mean_lr = np.divide(np.nansum(lr, 1), n_both, out=np.full(len(lr), np.nan), where=n_both > 0)
        enz_calls = {p: c for (e, p), c in calls.items() if e == enz}
        missing = set(enz_calls) - set(pid['Sequence'])
        n_rows = len(rows)
        for i, r in pid.iterrows():
            seq = r['Sequence']
            call = enz_calls.get(seq, ('', ''))[0]
            acc, s, how = ref.place(seq, r['Leading razor protein'], [p for p in r['Proteins'].split(';') if p],
                                    [g for g in r['Gene names'].split(';') if g])
            if acc is None:
                dropped[(how, call == 'RBDpep')] += 1
                if call == 'RBDpep':
                    s1_check['rbdpep_not_placed'] += 1
                continue
            how_count[how] += 1
            fa, fb = rbdmap_fragment(ref.seq[acc], site, s, s + len(seq))
            if call == 'RBDpep':  # Table S1's fragment, where the peptide is placed on Table S1's own entry
                _, frag, prot = enz_calls[seq]
                s1_check['rbdpep_on_another_entry' if prot != acc else
                         'rbdpep_fragment_same' if ref.seq[acc][fa:fb] == frag else 'rbdpep_fragment_differs'] += 1
            rows.append([acc, ref.gene.get(acc, ''), s + 1, s + len(seq), seq, enz, fa + 1, fb, how,
                         int(((bound[i] > 0) | (released[i] > 0)).sum()), int(n_both[i]), int(inp[i].sum()),
                         '' if np.isnan(mean_lr[i]) else f'{mean_lr[i]:.2f}',
                         'true' if call == 'RBDpep' else 'false', 'true' if call == 'CandidateRBDpep' else 'false'])
        per_enzyme[enz] = {'peptides_in_file': n_all, 'measured': len(pid), 'kept': len(rows) - n_rows, 'replicates': reps,
                           'quantified_bound_or_released': int((((bound > 0) | (released > 0)).sum(1) > 0).sum()),
                           'table_s1_rbdpep': sum(c[0] == 'RBDpep' for c in enz_calls.values()),
                           'table_s1_candidate': sum(c[0] == 'CandidateRBDpep' for c in enz_calls.values()),
                           'table_s1_not_in_peptide_table': len(missing),
                           'rbdpep_with_bound_and_released_in_a_replicate': int(sum(n_both[k] > 0 for k, p in enumerate(pid['Sequence'])
                                                                                     if enz_calls.get(p, ('',))[0] == 'RBDpep')),
                           'rbdpep_median_log2_bound_over_released': float(np.nanmedian([mean_lr[k] for k, p in enumerate(pid['Sequence'])
                                                                                          if enz_calls.get(p, ('',))[0] == 'RBDpep'])),
                           'other_median_log2_bound_over_released': float(np.nanmedian([mean_lr[k] for k, p in enumerate(pid['Sequence'])
                                                                                         if not enz_calls.get(p)]))}
        log(f'RBDmap {enz}: {json.dumps(per_enzyme[enz])}')
    rows.sort(key=lambda x: (x[0], x[2], x[3], x[4], x[5]))
    columns = ['uniprot_id', 'gene', 'start', 'end', 'sequence', 'experiment', 'fragment_start', 'fragment_end', 'mapped_by',
               'quantified_bound_or_released', 'quantified_bound_and_released', 'quantified_input', 'log2_bound_over_released',
               'rna_bound', 'candidate_rna_bound']
    out.mkdir(parents=True, exist_ok=True)
    path = out / RBDMAP_TABLE
    tsv(path, columns, rows)
    counts = placed_counts(rows, dropped, 'rna_bound')
    summary = {**counts, 'placed_by': dict(how_count), 'by_experiment': per_enzyme, 'table_s1_check': dict(s1_check),
               'kept_rna_bound': sum(r[13] == 'true' for r in rows), 'kept_candidate': sum(r[14] == 'true' for r in rows),
               'proteins': len({r[0] for r in rows}), 'proteins_rna_bound': len({r[0] for r in rows if r[13] == 'true'}),
               'pride_licence': project.get('license')}
    log(f'{RBDMAP_TABLE}: {len(rows):,} peptides on {summary["proteins"]:,} proteins, {summary["kept_rna_bound"]:,} RNA-bound '
        f'(RBDpep) on {summary["proteins_rna_bound"]:,}; {path.stat().st_size / 1e6:.2f} MB')

    raw = [[p.name, p.stat().st_size, file_hash(p), saved(p)] for p in [*paths.values(), s1, *meta]]
    raw_path = out / RBDMAP_RAW
    tsv(raw_path, ['path', 'bytes', 'md5', 'saved'], raw)
    day = saved(paths[RBDMAP_FILES[0]])
    n_rb = summary['kept_rna_bound']
    manifest = [
        {'file': f'{out.name}/{RBDMAP_TABLE}', 'bytes': path.stat().st_size, 'md5': file_hash(path), 'rows': len(rows),
         'redistribution': 'checksum_only',
         'content': f'RBDmap (HeLa, 254 nm UV; oligo(dT) capture, Lys-C or Arg-C, second capture into RNA-bound (F3) and released '
                    f'(F2) fragments, trypsin, dimethyl labels) human tryptic peptides identified in the Lys-C and Arg-C experiments, '
                    f'from the PRIDE MaxQuant peptide and quantification tables: each peptide on the canonical sequence (exact, unique '
                    f'match) with its protease, its Lys-C or Arg-C fragment (the in-silico digest that reproduces Table S1\'s fragments), '
                    f'in how many replicates it was quantified in the RNA-bound or released fraction, in both, and in the input, its '
                    f'mean log2 RNA-bound/released intensity ratio (not normalised), and the authors\' call from Table S1 (rna_bound: '
                    f'RBDpep, 1% FDR; candidate_rna_bound: CandidateRBDpep, 10% FDR). {n_rb:,} RBDpeps on '
                    f'{summary["proteins_rna_bound"]:,} proteins; the others are the measured background. The search did not allow '
                    f'methylarginine',
         'source': f'{PRIDE_FTP.format(path=RBDMAP_PATH, name="RBDmap_{LysC,ArgC}_{peptideID,QUAN}.txt")}; calls from {RBDMAP_S1_URL}',
         'version_or_date': f'PRIDE {RBDMAP_ACC} (public {project.get("publicationDate")}; MaxQuant 1.3.0.5), downloaded {day}',
         'citation': RBDMAP_CITATION,
         'licence': f'{project.get("license")} (PRIDE {RBDMAP_ACC}: no licence from the submitters); the calls are from Table S1 of the '
                    f'open access article (CC BY 4.0). Not redistributed',
         'derived_from': f'{", ".join(RBDMAP_FILES)}; {RBDMAP_S1} ({RBDMAP_RAW})',
         'how_made_or_checked': f'scripts/43f_stage_peptide_backgrounds.py (PRIDE files\' SHA-1 checked against PRIDE, Table S1\'s MD5 '
                                f'against raw_cache_files.tsv; {len(rows):,} peptides kept, {counts["dropped"]:,} dropped: '
                                + ', '.join(f'{v:,} {k.replace("_", " ")}' for k, v in counts['dropped_by_reason'].items())
                                + f'; {counts["dropped_rna_bound"]:,} of them RBDpeps). Every Table S1 RBDpep is in the peptide and '
                                  f'quantification tables; the fragment computed here equals Table S1\'s for '
                                  f'{s1_check["rbdpep_fragment_same"]:,} of the {s1_check["rbdpep_fragment_same"] + s1_check["rbdpep_fragment_differs"]:,} '
                                  f'RBDpeps placed on Table S1\'s own entry ({s1_check["rbdpep_on_another_entry"]:,} are placed on '
                                  f'another: Table S1 gives an unreviewed entry, none, or a paralog sharing the peptide)',
         'original_name': 'RBDmap_LysC_peptideID.txt; RBDmap_ArgC_peptideID.txt'},
        {'file': f'{out.name}/{RBDMAP_RAW}', 'bytes': raw_path.stat().st_size, 'md5': file_hash(raw_path), 'rows': len(raw),
         'redistribution': 'git',
         'content': f'The files behind {RBDMAP_TABLE}: the four PRIDE {RBDMAP_ACC} peptide and quantification tables, Table S1 '
                    f'(the RBDpep calls) and PRIDE\'s project and file listings (licence; SHA-1 of each file), with their size, '
                    f'MD5 and date',
         'source': f'{PRIDE_FTP.format(path=RBDMAP_PATH, name="RBDmap_{LysC,ArgC}_{peptideID,QUAN}.txt")}; {RBDMAP_S1_URL}; '
                   f'{PRIDE_API.format(acc=RBDMAP_ACC)}',
         'version_or_date': day, 'citation': '', 'licence': '', 'derived_from': '',
         'how_made_or_checked': 'scripts/43f_stage_peptide_backgrounds.py', 'original_name': ''}]
    return summary, manifest


# ---- preview: methylarginines against the other arginines of the same proteins ---------------------------------------------

TYPE_CLASSES = [  # as scripts/37's xl_types: each class against the unmethylated arginines (other methylarginines left out)
    ('SDMA_ip', lambda s: 'SDMA' in s['ip']),
    ('ADMA_ip', lambda s: 'ADMA' in s['ip'] and 'SDMA' not in s['ip']),
    ('MMA_ip_only', lambda s: s['ip'] == {'MMA'}),
    ('Rme2', lambda s: s['state'] in ('Rme2', 'Rme1_Rme2')),
    ('Rme1_only', lambda s: s['state'] == 'Rme1'),
]


def intervals(df: pd.DataFrame, a: str, b: str) -> dict:
    d = defaultdict(set)
    for acc, s, e in zip(df['uniprot_id'], df[a], df[b]):
        d[acc].add((int(s), int(e)))
    return {k: sorted(v) for k, v in d.items()}


def inside(p: int, iv) -> bool:
    return any(s <= p <= e for s, e in iv)


def preview_dataset(name: str, ref_seq: dict, methyl: dict, info: dict, sets: dict, tests: dict, classes: list) -> list:
    """sets: {key: {acc: [(start, end)]}} with 'region' (RNA-bound regions, the strata) and the others the tests use.
    tests: {test: (keep(acc, p) -> bool, inside(acc, p) -> bool)}; tests a_ and b_ are repeated for each methylation class."""
    strata = sorted(sets['region'])
    results = []
    for test, (keep, hit) in tests.items():
        for cls, member in [('all', None)] + classes:
            if cls != 'all' and not test.startswith(('a_', 'b_')):
                continue
            tables = []
            for acc in strata:
                seq, m = ref_seq[acc], methyl.get(acc, {})
                t = [0, 0, 0, 0]
                for i, aa in enumerate(seq):
                    p = i + 1
                    if aa != 'R' or not keep(acc, p):
                        continue
                    if p in m:
                        if member is not None and not member(info[(acc, p)]):
                            continue
                        t[0 if hit(acc, p) else 1] += 1
                    else:
                        t[2 if hit(acc, p) else 3] += 1
                tables.append(tuple(t))
            r = mantel_haenszel(tables)
            r.update(dataset=name, test=test, methylation=cls)
            results.append(r)
    return results


def preview(args) -> dict:
    ref_seq = Reference(Path(args.fasta), Path(args.entries), Path(args.history)).seq
    sites = pd.read_csv(args.sites, sep='\t', dtype=str, keep_default_na=False)
    has_types = {'maron2021_s3_antibodies', 'methyl_state'} <= set(sites.columns)
    methyl, info, off = defaultdict(set), {}, 0
    for r in sites.itertuples(index=False):
        acc, pos = r.uniprot_accession, int(r.position)
        s = ref_seq.get(acc, '')
        if pos <= len(s) and s[pos - 1] == 'R':
            methyl[acc].add(pos)
            info[(acc, pos)] = {'ip': {x for x in getattr(r, 'maron2021_s3_antibodies', '').split('; ') if x},
                                'state': getattr(r, 'methyl_state', '')}
        else:
            off += 1
    methyl = dict(methyl)
    classes = TYPE_CLASSES if has_types else []
    out = Path(args.out)
    results, coverage = [], {'methylarginines_in_release': sum(len(v) for v in methyl.values()),
                             'methylated_proteins_in_release': len(methyl), 'release_sites_not_R_on_reference': off}

    def cover(name, sets):
        strata = set(sets['region'])
        m_in = [(a, p) for a in sorted(strata & set(methyl)) for p in methyl[a]]
        coverage[name] = {'proteins_with_rna_bound_region': len(strata), 'methylated_proteins_with_rna_bound_region': len({a for a, _ in m_in}),
                          'methylarginines_in_those_proteins': len(m_in),
                          **{f'methylarginines_inside_{k}': sum(inside(p, v.get(a, [])) for a, p in m_in) for k, v in sets.items()}}

    if (out / CAPRI_TABLE).exists() and args.only in ('', 'capri'):
        t = pd.read_csv(out / CAPRI_TABLE, sep='\t', dtype=str, keep_default_na=False)
        adj = t[t['adjacent_peptide'] == 'true']
        sets = {'region': intervals(adj, 'lysc_start', 'lysc_end'), 'peptide': intervals(adj, 'start', 'end'),
                'measured_peptide': intervals(t, 'start', 'end'), 'measured_region': intervals(t, 'lysc_start', 'lysc_end')}
        g = lambda k: (lambda acc, p: inside(p, sets[k].get(acc, [])))  # noqa: E731
        tests = {'a_in_region_all_arginines': (lambda acc, p: True, g('region')),
                 'b_in_region_among_measured_peptides': (g('measured_peptide'), g('region')),
                 'c_in_region_among_measured_regions': (g('measured_region'), g('region')),
                 'x_in_adjacent_peptide_all_arginines': (lambda acc, p: True, g('peptide')),
                 'x_in_adjacent_peptide_among_measured_peptides': (g('measured_peptide'), g('peptide'))}
        cover('capri', sets)
        results += preview_dataset('capri', ref_seq, methyl, info, sets, tests, classes)
    if (out / RBDMAP_TABLE).exists() and args.only in ('', 'rbdmap'):
        t = pd.read_csv(out / RBDMAP_TABLE, sep='\t', dtype=str, keep_default_na=False)
        rb = t[t['rna_bound'] == 'true']
        q = t[t['quantified_bound_or_released'].astype(int) > 0]
        sets = {'region': intervals(rb, 'fragment_start', 'fragment_end'), 'peptide': intervals(rb, 'start', 'end'),
                'measured_peptide': intervals(q, 'start', 'end'), 'measured_region': intervals(q, 'fragment_start', 'fragment_end'),
                'identified_peptide': intervals(t, 'start', 'end'), 'identified_region': intervals(t, 'fragment_start', 'fragment_end')}
        g = lambda k: (lambda acc, p: inside(p, sets[k].get(acc, [])))  # noqa: E731
        tests = {'a_in_region_all_arginines': (lambda acc, p: True, g('region')),
                 'b_in_region_among_measured_peptides': (g('measured_peptide'), g('region')),
                 'c_in_region_among_measured_regions': (g('measured_region'), g('region')),
                 'x_in_region_among_identified_peptides_incl_input_only': (g('identified_peptide'), g('region')),
                 'x_in_region_among_identified_regions_incl_input_only': (g('identified_region'), g('region')),
                 'x_in_rbdpep_peptide_all_arginines': (lambda acc, p: True, g('peptide')),
                 'x_in_rbdpep_peptide_among_measured_peptides': (g('measured_peptide'), g('peptide'))}
        cover('rbdmap', sets)
        results += preview_dataset('rbdmap', ref_seq, methyl, info, sets, tests, classes)
    log(json.dumps(coverage, indent=1))
    log('dataset\ttest\tmethylation\ta\tb\tc\td\tOR\t95% CI\tp\tinformative proteins / proteins')
    for r in results:
        log(f'{r["dataset"]}\t{r["test"]}\t{r["methylation"]}\t{r["a"]}\t{r["b"]}\t{r["c"]}\t{r["d"]}\t{r["OR"]:.3f}\t'
            f'{r["ci_low"]:.3f}-{r["ci_high"]:.3f}\t{r["p"]:.3g}\t{r["informative_proteins"]} / {r["proteins"]}')
    res = {'coverage': coverage, 'tests': results, 'sites': str(args.sites)}
    if args.preview_out:
        Path(args.preview_out).write_text(json.dumps(res, indent=1) + '\n')
    return res


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0], formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--only', choices=['', 'capri', 'rbdmap'], default='', help='stage (and preview) one dataset')
    ap.add_argument('--cache', default=str(root / 'PTM_data' / 'external'), help='download folder (capri/, rbdmap/; kept between runs)')
    ap.add_argument('--out', default=str(root / 'inputs' / 'interactions'))
    ap.add_argument('--manifest', default=str(root / 'inputs' / 'MANIFEST.tsv'), help="'' to leave the manifest alone")
    ap.add_argument('--capri-adj', default=str(root / 'inputs' / 'interactions' / CAPRI_ADJ), help="the authors' ADJ-peptides (Supplementary Data 3)")
    ap.add_argument('--raw-cache', default=str(root / 'inputs' / 'interactions' / 'raw_cache_files.tsv'),
                    help='the MD5 of the RBDmap Table S1 scripts/37 read')
    ap.add_argument('--fasta', default=str(root / 'inputs' / 'reference' / 'uniprot_human_reviewed_2026_03.fasta'))
    ap.add_argument('--entries', default=str(root / 'inputs' / 'reference' / 'uniprot_human_reviewed_2026_03_entries.tsv'))
    ap.add_argument('--history', default=str(root / 'inputs' / 'reference' / 'uniprot_accession_history_2026_03.tsv'))
    ap.add_argument('--no-stage', action='store_true', help='only run the preview on existing tables')
    ap.add_argument('--preview', action='store_true', help='compare methylarginines with the other arginines (prints)')
    ap.add_argument('--sites', default=str(root / 'releases' / 'v1.2.0' / 'human_methylarginine_sites_v1.2.0.tsv'))
    ap.add_argument('--preview-out', default='', help='also write the preview as JSON here (outside inputs/)')
    args = ap.parse_args()
    if not args.no_stage:
        ref = Reference(Path(args.fasta), Path(args.entries), Path(args.history))
        manifest = []
        for name, stage in (('capri', stage_capri), ('rbdmap', stage_rbdmap)):
            if args.only in ('', name):
                summary, rows = stage(args, ref)
                log(json.dumps({name: summary}, indent=1))
                manifest += rows
        if args.manifest:
            write_manifest(Path(args.manifest), manifest)
    if args.preview:
        preview(args)


if __name__ == '__main__':
    main()
